"""读取登分表：自动找表头、识别科目和年级、统一班级与姓名写法、剔除成绩不全的学生。"""
from __future__ import annotations

import re
from dataclasses import dataclass, field

import openpyxl
import pandas as pd

from .config import Config, School, normalize_class


@dataclass
class ScoreData:
    df: pd.DataFrame                                  # 列：班级 姓名 考号 + 各科
    grade: str
    subjects: list[str]
    excluded: list[dict] = field(default_factory=list)   # 成绩不全、未计入的学生
    ignored_cols: list[str] = field(default_factory=list)  # 表头里没认出是科目的列
    empty_subjects: list[str] = field(default_factory=list)  # 有科目列但整列没有成绩，本次不计算
    sheet: str = ""

    @property
    def classes(self) -> list[str]:
        return sorted(self.df["班级"].unique(), key=lambda c: int(re.sub(r"\D", "", c)))


NON_SUBJECT_COLS = {"序号", "班级", "姓名", "考号", "学号", "总分", "总成绩", "平均分",
                    "名次", "班名次", "级名次", "排名", "备注", "性别", "班主任"}


def _find_header(wb, school: School):
    for ws in wb.worksheets:
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
            cells = [re.sub(r"\s", "", str(c)) if c is not None else "" for c in row]
            if "姓名" in cells and "班级" in cells:
                subj_cols = {school.canonical(c): j for j, c in enumerate(cells) if school.canonical(c)}
                if subj_cols:
                    ignored = [c for c in cells if c and c not in NON_SUBJECT_COLS
                               and not school.canonical(c)]
                    return ws, i, cells, subj_cols, ignored
    raise ValueError("登分表中找不到表头：需要有“姓名”“班级”两列，并至少有一门科目（科目名称见 学校设置.yaml）")


def _detect_grade(raw_classes, school: School) -> str | None:
    for raw in raw_classes:
        s = str(raw)
        for g, ch in school.grades.items():
            if ch in s:
                return g
    return None


def load_scores(path, school: School, grade: str | None = None) -> ScoreData:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, hr, header, subj_cols, ignored = _find_header(wb, school)
    subjects = [s for s in school.subjects if s in subj_cols]       # 按科目总表的顺序
    col = {name: j for j, name in enumerate(header) if name}
    id_col = next((c for c in ("考号", "学号") if c in col), None)
    raw_rows = [r for r in ws.iter_rows(min_row=hr + 1, values_only=True)
                if r[col["姓名"]] is not None and str(r[col["姓名"]]).strip()]
    grade = grade or _detect_grade([r[col["班级"]] for r in raw_rows], school)
    if not grade:
        raise ValueError("无法从班级名称判断年级，请在命令中用 --年级 指定（如 --年级 九年级）")
    gc = school.grades[grade]
    # 整列一个分数都没有的科目 = 本次没考，不计算（不能因为它把学生判成“成绩不全”）
    empty = [s for s in subjects
             if not any(isinstance(r[subj_cols[s]], (int, float)) for r in raw_rows)]
    subjects = [s for s in subjects if s not in empty]
    if not subjects:
        raise ValueError("登分表里所有科目列都没有成绩，请检查是否填好了分数。")
    rows, excluded = [], []
    for r in raw_rows:
        name = re.sub(r"\s", "", str(r[col["姓名"]]))
        cls = normalize_class(r[col["班级"]], gc)
        sid = str(r[col[id_col]]).strip() if id_col and r[col[id_col]] is not None else ""
        scores = [r[subj_cols[s]] for s in subjects]
        if not all(isinstance(v, (int, float)) for v in scores):
            reason = "；".join(f"{s}：{v if v not in (None, '') else '空'}"
                              for s, v in zip(subjects, scores) if not isinstance(v, (int, float)))
            excluded.append({"班级": cls, "姓名": name, "原因": reason})
            continue
        rows.append([cls, name, sid] + [float(v) for v in scores])
    df = pd.DataFrame(rows, columns=["班级", "姓名", "考号"] + subjects)
    df["_序"] = df["班级"].str.replace(r"\D", "", regex=True).astype(int)
    df = df.sort_values(["_序", "考号"], kind="stable").drop(columns="_序").reset_index(drop=True)
    return ScoreData(df=df, grade=grade, subjects=subjects, excluded=excluded,
                     ignored_cols=ignored, empty_subjects=empty, sheet=ws.title)


def check_data(data: ScoreData, cfg: Config) -> list[str]:
    """检查明显问题：分数超过满分、负分、重名。返回提示列表（不阻止计算，除超满分外）。"""
    msgs = []
    for s in cfg.subjects:
        over = data.df[data.df[s] > cfg.full[s]]
        if len(over):
            msgs.append(f"【错误】{s}有 {len(over)} 人分数超过满分 {cfg.full[s]:g}（最高 {over[s].max():g}），"
                        f"请检查满分设置或登分是否有误")
        if (data.df[s] < 0).any():
            msgs.append(f"【错误】{s}有负分")
    if data.empty_subjects:
        msgs.append(f"【提示】{'、'.join(data.empty_subjects)} 整列没有成绩，本次不计算。")
    if data.ignored_cols:
        msgs.append(f"【提示】登分表里这些列没当成科目，已忽略：{'、'.join(data.ignored_cols)}。"
                    f"如果其中有本次要统计的科目，请把列名改成标准科目名，"
                    f"或在 config/学校设置.yaml 的“科目别名”里加上这种写法。")
    dup = data.df[data.df.duplicated("姓名", keep=False)]
    if len(dup):
        names = "、".join(f"{n}（{'、'.join(g['班级'])}）" for n, g in dup.groupby("姓名"))
        msgs.append(f"【提示】有重名学生，已按班级区分：{names}")
    return msgs
