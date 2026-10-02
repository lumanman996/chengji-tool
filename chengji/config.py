"""学校设置、任课总表的读取，以及每次计算用的配置（Config）。"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import yaml

CN_DIGITS = "零一二三四五六七八九"


@dataclass
class Teacher:
    subject: str
    name: str
    classes: list[str]


@dataclass
class School:
    name: str
    org_template: str
    grades: dict[str, str]                 # 七年级 -> 七
    full: dict[str, float]                 # 科目 -> 满分（0 = 待确认）
    grade_full: dict[str, dict[str, float]]
    aliases: dict[str, str]
    schemes: dict[str, dict]               # 算法方案名 -> {优秀率, 及格率, 平均分, 平均分折算百分制}
    structures: dict[str, dict] = field(default_factory=dict)   # 结构分方案名 -> {项目: 分值}
    pdf_sections: dict[str, bool] = field(default_factory=dict) # 发布版内容 -> 是否导出（没写的默认导出）
    grade_subjects: dict[str, list[str]] = field(default_factory=dict)   # 七年级/八年级/九年级/中考 -> 固定科目

    def expected_subjects(self, grade: str, scheme: str) -> list[str]:
        """这个年级、这个方案按规定应该有哪些科目（中考方案看“中考”那一行）。没规定返回空。"""
        return list(self.grade_subjects.get(scheme) or self.grade_subjects.get(grade) or [])
    pass_ratio: float = 0.6
    excellent_ratio: float = 0.8
    promote_ratio: float = 0.85
    near_range: float = 15
    high_ratio: float = 0.75
    top_n: list[int] = field(default_factory=lambda: [10, 20, 40, 50, 80, 100, 150, 200])

    @property
    def subjects(self) -> list[str]:
        return list(self.full)

    def canonical(self, header: str) -> str | None:
        h = re.sub(r"\s", "", str(header or ""))
        h = self.aliases.get(h, h)
        return h if h in self.full else None

    def grade_of_char(self, ch: str) -> str | None:
        return next((g for g, c in self.grades.items() if c == ch), None)


def _yes(v) -> bool:
    return str(v).strip().lower() in ("是", "true", "yes", "1", "对")


def load_school(path: str | Path, local: str | Path | None = None) -> School:
    """读 学校设置.yaml。local = 只在本机生效的设置文件（如 data/本校设置.yaml），有的话盖在上面。"""
    d = yaml.safe_load(Path(path).read_text(encoding="utf-8"))
    if local and Path(local).is_file():
        d.update(yaml.safe_load(Path(local).read_text(encoding="utf-8")) or {})
    schemes = {}
    for name, w in (d.get("算法方案") or {}).items():
        schemes[name] = {"优秀率": float(w.get("优秀率", 0)), "及格率": float(w.get("及格率", 0)),
                         "平均分": float(w.get("平均分", 0)), "折算": _yes(w.get("平均分折算百分制", "是"))}
    if not schemes:      # 兼容旧写法“权重”
        w = d.get("权重") or {}
        schemes["平时"] = {"优秀率": float(w.get("优秀率", 0.25)), "及格率": float(w.get("及格率", 0.25)),
                         "平均分": float(w.get("平均分", 0.5)), "折算": True}
    return School(
        name=d.get("学校", ""),
        org_template=d.get("落款", "{年级}组"),
        grades=dict(d["年级"]),
        full={k: float(v or 0) for k, v in d["科目满分"].items()},
        grade_full={g: {k: float(v) for k, v in (m or {}).items()} for g, m in (d.get("年级满分调整") or {}).items()},
        aliases=dict(d.get("科目别名") or {}),
        schemes=schemes,
        structures={k: {i: float(v or 0) for i, v in (m or {}).items()}
                    for k, m in (d.get("结构分方案") or {}).items()},
        pdf_sections={k: _yes(v) for k, v in (d.get("发布版内容") or {}).items()},
        grade_subjects={k: [str(x) for x in (v or [])] for k, v in (d.get("年级科目") or {}).items()},
        pass_ratio=float(d.get("及格线比例", 0.6)),
        excellent_ratio=float(d.get("优秀线比例", 0.8)),
        promote_ratio=float(d.get("升学线比例", 0.85)),
        near_range=float(d.get("临界范围", 15)),
        high_ratio=float(d.get("高分线比例", 0.75)),
        top_n=list(d.get("前N名", [10, 20, 40, 50, 80, 100, 150, 200])),
    )


@dataclass
class Config:
    """一次计算（某年级、某次考试）用的全部设置。"""
    school: str
    grade: str          # 九年级
    grade_char: str     # 九
    exam: str
    date: str
    org: str            # 落款
    subjects: list[str]
    full: dict[str, float]
    pass_ratio: float = 0.6
    excellent_ratio: float = 0.8
    w_excellent: float = 0.25
    w_pass: float = 0.25
    w_avg: float = 0.5
    normalize: bool = True          # 平均分是否折算百分制
    scheme: str = "平时"
    promote_ratio: float = 0.85
    near_range: float = 15
    high_line: float = 0
    band_width: float = 0
    top_n: list[int] = field(default_factory=lambda: [10, 20, 40, 50, 80, 100, 150, 200])
    teachers: list[Teacher] = field(default_factory=list)
    structure: dict[str, float] = field(default_factory=dict)   # 结构分各项分值（空 = 不算结构分）
    enrolled: dict[str, int] = field(default_factory=dict)      # 班级 -> 应考人数
    prev_rank: dict[str, int] = field(default_factory=dict)     # 班级 -> 上次考试结构分综合名次

    @property
    def total_full(self) -> float:
        return sum(self.full.values())

    def pass_line(self, s: str) -> float:
        return self.full[s] * self.pass_ratio

    def excellent_line(self, s: str) -> float:
        return self.full[s] * self.excellent_ratio

    def avg_term(self, avg: float, s: str) -> float:
        """平均分在得分公式中的取值：折算百分制或原始分。"""
        return avg / self.full[s] * 100 if self.normalize else avg

    def formula_text(self) -> str:
        parts = []
        if self.w_excellent:
            parts.append(f"优秀率×{self.w_excellent:.0%}")
        if self.w_pass:
            parts.append(f"{'合格率' if self.scheme == '中考' else '及格率'}×{self.w_pass:.0%}")
        if self.w_avg:
            parts.append(f"平均分（{'百分制' if self.normalize else '原始分'}）×{self.w_avg:.0%}")
        return " + ".join(parts)


def make_config(school: School, grade: str, subjects: list[str], exam: str, date: str = "",
                teachers: list[Teacher] | None = None, scheme: str = "平时",
                promote_ratio: float | None = None, enrolled: dict[str, int] | None = None,
                prev_rank: dict[str, int] | None = None) -> Config:
    if scheme not in school.schemes:
        raise ValueError(f"没有名为“{scheme}”的算法方案，可选：{'、'.join(school.schemes)}")
    w = school.schemes[scheme]
    pr = school.promote_ratio if promote_ratio is None else promote_ratio
    if not 0 < pr <= 1:
        raise ValueError(f"升学线比例应在 0～100% 之间，现在是 {pr}")
    full = {s: school.grade_full.get(grade, {}).get(s, school.full[s]) for s in subjects}
    missing = [s for s, v in full.items() if not v]
    if missing:
        raise ValueError(f"以下科目的满分还没有设置：{'、'.join(missing)}。请在 config/学校设置.yaml 的“科目满分”中填写。")
    tf = sum(full.values())
    return Config(
        school=school.name, grade=grade, grade_char=school.grades[grade], exam=exam, date=date,
        org=school.org_template.replace("{年级}", grade), subjects=subjects, full=full,
        pass_ratio=school.pass_ratio, excellent_ratio=school.excellent_ratio,
        w_excellent=w["优秀率"], w_pass=w["及格率"], w_avg=w["平均分"], normalize=w["折算"], scheme=scheme,
        promote_ratio=pr, near_range=school.near_range,
        high_line=round(tf * school.high_ratio / 10) * 10,
        band_width=20 if tf <= 400 else 50,
        top_n=school.top_n,
        teachers=[t for t in (teachers or []) if t.subject in subjects],
        structure={k: v for k, v in school.structures.get(scheme, {}).items() if v},
        enrolled=dict(enrolled or {}),
        prev_rank=dict(prev_rank or {}),
    )


# ---------------- 班级、任课 ----------------
def normalize_class(raw, grade_char: str) -> str:
    """把“九（1）”“九1”“91”“九年级1班”等写法统一成“九1”。"""
    s = re.sub(r"[\s（）()班级年]", "", str(raw).strip())
    digits = re.findall(r"\d+", s)
    if not digits:
        raise ValueError(f"无法识别班级：{raw!r}")
    num = digits[-1]
    gnum = str(CN_DIGITS.index(grade_char)) if grade_char in CN_DIGITS else None
    if s.isdigit() and gnum and len(num) >= 2 and num[0] == gnum:   # “91” = 九1
        num = num[1:]
    return f"{grade_char}{int(num)}"


def parse_classes(text, grade_char: str) -> list[str]:
    """解析任教班级：“9192”“939495”“九1、九2”“七1 七2”“93”。"""
    parts = [p for p in re.split(r"[\s、，,；;/]+", str(text).strip()) if p]
    gnum = str(CN_DIGITS.index(grade_char)) if grade_char in CN_DIGITS else None
    out: list[str] = []
    for code in parts:
        if code.isdigit() and gnum and len(code) % 2 == 0 and all(code[i] == gnum for i in range(0, len(code), 2)):
            out += [f"{grade_char}{int(code[i + 1])}" for i in range(0, len(code), 2)]
        else:
            out.append(normalize_class(code, grade_char))
    return out


def load_teachers(path: str | Path, school: School, grade: str) -> list[Teacher]:
    """从任课总表读取某个年级的任课安排（表头：年级、学科、教师、任教班级）。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["任课总表"] if "任课总表" in wb.sheetnames else wb.worksheets[0]
    hr = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        cells = [str(c).strip() if c else "" for c in row]
        if all(k in cells for k in ("年级", "学科", "教师", "任教班级")):
            hr, idx = i, {k: cells.index(k) for k in ("年级", "学科", "教师", "任教班级")}
            break
    if not hr:
        raise ValueError("任课总表中找不到表头（年级、学科、教师、任教班级）")
    gc = school.grades[grade]
    out = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        g, subj, name, cls = (row[idx[k]] for k in ("年级", "学科", "教师", "任教班级"))
        if not (g and subj and name and cls):
            continue
        g = str(g).strip()
        g = g if g.endswith("年级") else (school.grade_of_char(g[0]) or g)
        if g != grade:
            continue
        subj = school.canonical(subj) or str(subj).strip()
        out.append(Teacher(subj, re.sub(r"\s", "", str(name)), parse_classes(cls, gc)))
    return out


def load_enrolled(path: str | Path, school: School, grade: str) -> dict[str, int]:
    """从任课总表的「班级信息」工作表读取某个年级各班的应考人数（表头：年级、班级、应考人数）。
    没有这张工作表时返回空字典。"""
    wb = openpyxl.load_workbook(path, data_only=True)
    if "班级信息" not in wb.sheetnames:
        return {}
    ws = wb["班级信息"]
    keys = ("年级", "班级", "应考人数")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        cells = [str(c).strip() if c else "" for c in row]
        cells = ["应考人数" if c == "学籍人数" else c for c in cells]      # 兼容旧叫法
        if all(k in cells for k in keys):
            hr, idx = i, {k: cells.index(k) for k in keys}
            break
    else:
        raise ValueError("「班级信息」工作表中找不到表头（年级、班级、应考人数）")
    gc = school.grades[grade]
    out = {}
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        g, cls, n = (row[idx[k]] for k in keys)
        if not (g and cls) or not isinstance(n, (int, float)):
            continue
        g = str(g).strip()
        g = g if g.endswith("年级") else (school.grade_of_char(g[0]) or g)
        if g == grade:
            out[normalize_class(cls, gc)] = int(n)
    return out
