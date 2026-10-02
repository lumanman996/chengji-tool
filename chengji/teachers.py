"""任课总表的读和写（给界面里的“任课总表编辑”用）。

文件还是那个 Excel（data/任课总表.xlsx），格式和模板一样：
  「任课总表」  年级、学科、教师、任教班级
  「班级信息」  年级、班级、应考人数
所以在界面里改和用 Excel 改是一回事，两边可以混着用。
"""
from __future__ import annotations

import re
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .config import School, normalize_class, parse_classes

F = "微软雅黑"
_HEAD = PatternFill("solid", fgColor="DDEBF7")
_THIN = Border(*[Side(style="thin", color="BFBFBF")] * 4)

TIPS = [
    "任课总表填写说明",
    "",
    "1. 第 2 行是表头，不要改动文字；从第 3 行开始，一位教师教一个学科占一行。",
    "2. 年级：写 七年级 / 八年级 / 九年级（只写 七 / 八 / 九 也行）。",
    "3. 学科：语文、数学、英语、物理、化学、道法、历史、生物、地理、体育；自己加的科目也可以写。",
    "4. 任教班级：年级用汉字，班号用数字，多个班用顿号隔开，例如 九1、九2。",
    "   不要写成 9192、91 这样的纯数字。",
    "5. 一位教师教两个学科，就写两行；教两个年级，也写两行。",
    "6. 全校一张表就够了。算成绩时，程序自动挑出这次考试的年级和科目对应的教师，",
    "   表里多出来的年级、科目不影响计算。",
    "7. 教师姓名不要带空格；同名教师靠“学科 + 任教班级”区分。",
    "",
    "「班级信息」工作表（算班级结构分的参考率用）：",
    "8. 每个班一行：年级、班级、应考人数。班级写法同上，如 九1。",
    "9. 应考人数一学期填一次，有转入转出时改一下。",
    "10. 应考人数可以不填：没填的班按实考人数算（参考率按 100%）。",
    "",
    "这张表也可以在软件的“设置 → 任课总表 → 在这里编辑”里直接修改。",
]


def _grade(raw, school: School) -> str:
    g = str(raw or "").strip()
    if not g:
        return ""
    if g in school.grades:
        return g
    return school.grade_of_char(g[0]) or g


def _find(ws, keys):
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        cells = [str(c).strip() if c else "" for c in row]
        cells = ["应考人数" if c == "学籍人数" else c for c in cells]
        if all(k in cells for k in keys):
            return i, {k: cells.index(k) for k in keys}
    return None, None


def read_table(path, school: School) -> dict:
    """读出任课安排和班级信息（没填全的行也读出来，方便在界面里接着填）。文件不存在返回空表。"""
    out = {"rows": [], "classes": []}
    if not Path(path).is_file():
        return out
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["任课总表"] if "任课总表" in wb.sheetnames else wb.worksheets[0]
    hr, idx = _find(ws, ("年级", "学科", "教师", "任教班级"))
    if hr:
        for row in ws.iter_rows(min_row=hr + 1, values_only=True):
            g, sub, name, cls = (row[idx[k]] if idx[k] < len(row) else None for k in ("年级", "学科", "教师", "任教班级"))
            g = _grade(g, school)
            if g not in school.grades or not (sub or name or cls):
                continue                                      # 说明文字、空行
            text = str(cls).strip() if cls not in (None, "") else ""
            if text:
                try:
                    text = "、".join(parse_classes(text, school.grades[g]))
                except ValueError:
                    pass
            out["rows"].append({"grade": g, "subject": school.canonical(sub) or str(sub or "").strip(),
                                "teacher": re.sub(r"\s", "", str(name or "")), "classes": text})
    if "班级信息" in wb.sheetnames:
        ci = wb["班级信息"]
        hr, idx = _find(ci, ("年级", "班级", "应考人数"))
        if hr:
            for row in ci.iter_rows(min_row=hr + 1, values_only=True):
                g, cls, n = (row[idx[k]] if idx[k] < len(row) else None for k in ("年级", "班级", "应考人数"))
                g = _grade(g, school)
                if g not in school.grades or not cls:
                    continue
                try:
                    cls = normalize_class(cls, school.grades[g])
                except ValueError:
                    continue
                out["classes"].append({"grade": g, "cls": cls, "enrolled": int(n) if isinstance(n, (int, float)) else ""})
    return out


def check_table(data: dict, school: School) -> dict:
    """检查并整理界面交来的任课总表。有问题抛 ValueError（说清楚是哪个年级、哪一行）。"""
    rows, classes = [], []
    for k, r in enumerate(data.get("rows") or [], start=1):
        g, sub = str(r.get("grade") or "").strip(), str(r.get("subject") or "").strip()
        name, cls = re.sub(r"\s", "", str(r.get("teacher") or "")), str(r.get("classes") or "").strip()
        if not (sub or name or cls):
            continue
        where = f"{g or '任课安排'}第 {k} 行"
        if g not in school.grades:
            raise ValueError(f"{where}：年级不对。")
        if not sub:
            raise ValueError(f"{where}：还没有填学科。")
        if (name and not cls) or (cls and not name):
            raise ValueError(f"{g} {sub}：教师和任教班级要一起填（{name or cls}）。")
        if cls:
            try:
                cls = "、".join(parse_classes(cls, school.grades[g]))
            except ValueError:
                raise ValueError(f"{g} {sub} {name}：任教班级“{cls}”看不懂，请写成 {school.grades[g]}1、{school.grades[g]}2 这样。") from None
        rows.append({"grade": g, "subject": school.canonical(sub) or sub, "teacher": name, "classes": cls})
    seen = set()
    for r in data.get("classes") or []:
        g, cls, n = str(r.get("grade") or "").strip(), str(r.get("cls") or "").strip(), r.get("enrolled")
        if not cls:
            continue
        if g not in school.grades:
            raise ValueError("班级信息里有一行的年级不对。")
        try:
            cls = normalize_class(cls, school.grades[g])
        except ValueError:
            raise ValueError(f"{g} 的班级“{cls}”看不懂，请写成 {school.grades[g]}1 这样。") from None
        if (g, cls) in seen:
            raise ValueError(f"班级信息里 {cls} 写了两次。")
        seen.add((g, cls))
        if n in (None, ""):
            n = ""
        else:
            try:
                n = int(float(n))
            except (TypeError, ValueError):
                raise ValueError(f"{cls} 的应考人数要填数字。") from None
            if n <= 0:
                raise ValueError(f"{cls} 的应考人数要大于 0。")
        classes.append({"grade": g, "cls": cls, "enrolled": n})
    order = list(school.grades)
    rows.sort(key=lambda r: order.index(r["grade"]))                    # 同年级内保持填写的顺序
    classes.sort(key=lambda r: (order.index(r["grade"]), int(re.sub(r"\D", "", r["cls"]) or 0)))
    return {"rows": rows, "classes": classes}


def write_table(path, data: dict, title: str = "任课总表"):
    """按模板的格式写成 Excel。data 要先经过 check_table。"""
    wb = openpyxl.Workbook()

    def sheet(ws, name, heads, lines, widths):
        ws.title = name
        ws["A1"] = title if name == "任课总表" else "班级信息（算参考率用：参考率 = 实考人数 ÷ 应考人数）"
        ws["A1"].font = Font(name=F, size=14, bold=True)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(heads))
        for j, h in enumerate(heads, 1):
            c = ws.cell(2, j, h)
            c.font, c.fill, c.border = Font(name=F, bold=True), _HEAD, _THIN
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, line in enumerate(lines, start=3):
            for j, v in enumerate(line, 1):
                c = ws.cell(i, j, v if v != "" else None)
                c.font, c.border = Font(name=F), _THIN
                if not (name == "任课总表" and j == 4):
                    c.alignment = Alignment(horizontal="center")
        for col, w in zip("ABCD", widths):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "A3"

    sheet(wb.active, "任课总表", ["年级", "学科", "教师", "任教班级"],
          [(r["grade"], r["subject"], r["teacher"], r["classes"]) for r in data["rows"]], [10, 10, 12, 26])
    sheet(wb.create_sheet("班级信息"), "班级信息", ["年级", "班级", "应考人数"],
          [(r["grade"], r["cls"], r["enrolled"]) for r in data["classes"]], [10, 10, 12])
    s = wb.create_sheet("填写说明")
    for i, t in enumerate(TIPS, 1):
        s.cell(i, 1, t).font = Font(name=F, bold=(i == 1), size=13 if i == 1 else 11)
    s.column_dimensions["A"].width = 86
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
