"""生成统计用 Excel（带公式，改参数后在 Excel/WPS 中自动重算）。

表：各班统计、班级结构分、教师排名、前85%分数段、临界生名单、参数设置、成绩明细、任课安排；
选了“和上次比”时再加：和上次比、名次进退、上次临界生（上次的数字是当时的结果，本次的数字是公式）。
"""
from __future__ import annotations

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from .analysis import RATE_ITEMS, TWO_RATES, Result

FONT = "微软雅黑"
_thin = Side(style="thin", color="999999")
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
HEAD = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
SUM = PatternFill("solid", fgColor="E2EFDA")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)


def sty(c, bold=False, fill=None, color="000000"):
    c.font = Font(name=FONT, size=11, bold=bold, color=color)
    c.alignment = CENTER
    c.border = BORDER
    if fill:
        c.fill = fill


def title(ws, text, ncol, size=16):
    ws["A1"] = text
    ws["A1"].font = Font(name=FONT, size=size, bold=True)
    ws["A1"].alignment = CENTER
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)


def note(ws, r, text, ncol, size=10):
    c = ws.cell(r, 1, text)
    c.font = Font(name=FONT, size=size, color="666666")
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)


def build_excel(R: Result, path, conclusions: list[str], compare: dict | None = None):
    cfg = R.cfg
    SUBS, NS, CL = cfg.subjects, len(cfg.subjects), R.classes
    nc = len(CL)
    wb = openpyxl.Workbook()

    # ---------- 参数设置
    p = wb.active
    p.title = "参数设置"
    title(p, "参数设置（黄色格子可修改，其他表会自动重新计算）", 5, 14)
    for j, h in enumerate(["学科", "满分", "及格线", "优秀线", "说明"], 1):
        sty(p.cell(3, j, h), True, HEAD)
    r_pass, r_exc = 4 + NS + 2, 4 + NS + 3
    PR = {}
    for i, s in enumerate(SUBS):
        r = 4 + i
        PR[s] = r
        sty(p.cell(r, 1, s), True)
        sty(p.cell(r, 2, cfg.full[s]), fill=INPUT, color="0000FF")
        sty(p.cell(r, 3, f"=B{r}*$B${r_pass}"))
        sty(p.cell(r, 4, f"=B{r}*$B${r_exc}"))
        sty(p.cell(r, 5, ""))
    rT = 4 + NS
    sty(p.cell(rT, 1, f"{NS}科总分"), True, SUM)
    sty(p.cell(rT, 2, f"=SUM(B4:B{rT - 1})"), True, SUM)
    for j in range(3, 6):
        sty(p.cell(rT, j, ""), fill=SUM)
    p.cell(rT + 1, 1, "比例与权重").font = Font(name=FONT, size=12, bold=True)
    items = [("及格线占满分比例", cfg.pass_ratio, "满分的百分之几算及格"),
             ("优秀线占满分比例", cfg.excellent_ratio, "满分的百分之几算优秀"),
             ("优秀率权重", cfg.w_excellent, "单科得分中优秀率的比重"),
             ("及格率权重", cfg.w_pass, "单科得分中及格率的比重"),
             ("平均分权重", cfg.w_avg, "单科得分中平均分的比重"),
             ("平均分折算百分制（1=是，0=否）", 1 if cfg.normalize else 0, "1：平均分÷满分×100 后计分（平时方案）；0：直接用原始平均分（中考方案）"),
             ("升学线比例（前百分之几）", cfg.promote_ratio, "按高中录取率，取全级总分前百分之几"),
             ("临界生范围（分数线上下各几分）", cfg.near_range, "分数线上下多少分以内算临界生")]
    W = cfg.structure if R.structure else {}
    sdesc = {"平均成绩": "本班总分平均分 ÷ 总分满分 × 分值", "全科合格率": "每科都合格的人数 ÷ 实考人数 × 分值",
             "全科优秀率": "每科都优秀的人数 ÷ 实考人数 × 分值", "参考率": "实考人数 ÷ 应考人数 × 分值",
             "进线率": "升学线以上人数 ÷ 实考人数 × 分值", "增值评价": "两率一分（平均成绩+全科合格率+全科优秀率）第1名得满分",
             "增值名次差": "两率一分每低一个名次，增值评价减几分", "增值进退步": "比上次的两率一分名次每进步/退步一名，加/减几分（不超过满分）",
             "前10名每人加分": "全级总分前10名（同分计入）每有一人，本班加几分"}
    for k, v in W.items():
        items.append((f"结构分：{k}" + ("（分值）" if k in RATE_ITEMS or k == "增值评价" else ""), v,
                      sdesc.get(k, ""), "0.00" if v < 1 else "0"))
    P = {}
    for i, (k, v, dsc, *fm) in enumerate(items):
        r = rT + 2 + i
        P[k] = f"参数设置!$B${r}"
        sty(p.cell(r, 1, k), True)
        c = p.cell(r, 2, v)
        sty(c, fill=INPUT, color="0000FF")
        c.number_format = fm[0] if fm else ("0%" if (v < 1 and "折算" not in k) else "0")
        sty(p.cell(r, 3, dsc))
        p.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    assert P["及格线占满分比例"].endswith(f"${r_pass}") and P["优秀线占满分比例"].endswith(f"${r_exc}")
    rn = rT + 2 + len(items) + 1
    p.cell(rn - 1, 1, f"当前算法方案：{cfg.scheme}　单科得分 = {cfg.formula_text()}").font = Font(name=FONT, size=11, bold=True, color="C00000")
    p.cell(rn, 1, "计算方法").font = Font(name=FONT, size=12, bold=True)
    notes = ["1. 及格率 = 成绩 ≥ 及格线的人数 ÷ 实考人数；优秀率 = 成绩 ≥ 优秀线的人数 ÷ 实考人数。",
             "2. 单科得分 = 优秀率×100×优秀率权重 + 及格率×100×及格率权重 + 平均分×平均分权重（权重为 0 的项不计）。",
             "3. 平均分折算百分制时 = 平均分 ÷ 满分 × 100（平时方案）；不折算时直接用原始平均分（中考方案）。",
             f"4. 班级综合得分 = {NS}科得分之和，按从高到低排名次。",
             "5. 成绩不全（缺考、请假、转学、空白）的学生不参与统计；0 分为真实成绩，参与统计；重名学生按班级区分。",
             "6. 应考人数 = 应该参加考试的在籍学生；实考人数 = 各科成绩齐全、计入统计的学生。"
             "各表的平均分、及格率、优秀率都按实考人数算。"]
    if W:
        notes.append(f"{len(notes) + 1}. 班级结构分（{cfg.scheme}）：各项直接按比例折分（比率 × 分值），"
                     "全科合格 / 全科优秀 = 每一科都达到合格线 / 优秀线；分值在上面“结构分”各行修改。")
    if R.data.excluded:
        notes.append(f"{len(notes) + 1}. 本次未计入：" + "；".join(f"{e['班级']} {e['姓名']}（{e['原因']}）" for e in R.data.excluded))
    for i, n in enumerate(notes):
        note(p, rn + 1 + i, n, 5, 11)
    for col, w in zip("ABCDE", [32, 10, 10, 10, 40]):
        p.column_dimensions[col].width = w

    # ---------- 成绩明细
    # 列：序号 | 班级 | 姓名 | 考号 | 各科… | 总分 | 级名次 | 班名次 | （全科合格 | 全科优秀）
    d = wb.create_sheet("成绩明细")
    heads = ["序号", "班级", "姓名", "考号"] + SUBS + ["总分", "级名次", "班名次"] + (["全科合格", "全科优秀"] if W else [])
    for j, h in enumerate(heads, 1):
        sty(d.cell(1, j, h), True, HEAD)
    N = 1 + len(R.data.df)
    SC = {s: L(5 + i) for i, s in enumerate(SUBS)}
    TC = L(5 + NS)
    for i, row in enumerate(R.data.df.itertuples(index=False), start=2):
        # 动态序号：只数筛选后看得见的行，所以筛选出某个班时仍是 1、2、3…
        # （末尾 *1 是为了不让 Excel 把最后一行当成“汇总行”而漏筛）
        sty(d.cell(i, 1, f"=SUBTOTAL(103,$B$2:B{i})*1"))
        for j, v in enumerate(row, 2):
            sty(d.cell(i, j, v))
        sty(d.cell(i, 5 + NS, f"=SUM(E{i}:{L(4 + NS)}{i})"))
        sty(d.cell(i, 6 + NS, f"=RANK({TC}{i},${TC}$2:${TC}${N})"))
        sty(d.cell(i, 7 + NS, f'=COUNTIFS($B$2:$B${N},$B{i},${TC}$2:${TC}${N},">"&{TC}{i})+1'))
        if W:
            for j, lc in ((8 + NS, "C"), (9 + NS, "D")):     # 参数设置 C 列 = 及格线，D 列 = 优秀线
                conds = ",".join(f"{SC[sub]}{i}>=参数设置!${lc}${PR[sub]}" for sub in SUBS)
                sty(d.cell(i, j, f"=IF(AND({conds}),1,0)"))
    d.column_dimensions["A"].width = 6
    for j in range(2, len(heads) + 1):
        d.column_dimensions[L(j)].width = 9
    d.freeze_panes = "E2"
    d.auto_filter.ref = f"A1:{L(len(heads))}{N}"        # 表头带筛选按钮
    rng = lambda col: f"成绩明细!${col}$2:${col}${N}"
    cls, tot = rng("B"), rng(TC)

    NORM = P["平均分折算百分制（1=是，0=否）"]

    def score_formula(r, A_, P_, E_, pr):
        return (f"={E_}{r}*100*{P['优秀率权重']}+{P_}{r}*100*{P['及格率权重']}"
                f"+{A_}{r}*IF({NORM}=1,100/参数设置!$B${pr},1)*{P['平均分权重']}")

    # ---------- 各班统计
    s = wb.create_sheet("各班统计", 0)
    ncol = 2 + 5 * NS + 2
    title(s, f"{cfg.exam} 各班成绩综合统计（{cfg.scheme}方案：{cfg.formula_text()}）", ncol)
    sty(s.cell(2, 1, "班级"), True, HEAD); s.merge_cells("A2:A3")
    sty(s.cell(2, 2, "实考人数"), True, HEAD); s.merge_cells("B2:B3")
    for k, sub in enumerate(SUBS):
        c0 = 3 + 5 * k
        sty(s.cell(2, c0, f"{sub}（{cfg.full[sub]:g}）"), True, HEAD)
        s.merge_cells(start_row=2, start_column=c0, end_row=2, end_column=c0 + 4)
        for j, h in enumerate(["总分", "平均分", "及格率", "优秀率", "得分"]):
            sty(s.cell(3, c0 + j, h), True, HEAD)
    cC = 3 + 5 * NS
    sty(s.cell(2, cC, f"{NS}科综合得分"), True, SUM); s.merge_cells(start_row=2, start_column=cC, end_row=3, end_column=cC)
    sty(s.cell(2, cC + 1, "名次"), True, SUM); s.merge_cells(start_row=2, start_column=cC + 1, end_row=3, end_column=cC + 1)
    for r in (2, 3):
        for c in range(1, ncol + 1):
            s.cell(r, c).border = BORDER
    for i, cn in enumerate(CL + ["全年级"]):
        r = 4 + i
        g = cn == "全年级"
        fl = SUM if g else None
        sty(s.cell(r, 1, cn), True, fl)
        sty(s.cell(r, 2, f"=COUNTA({cls})" if g else f"=COUNTIF({cls},A{r})"), fill=fl)
        comps = []
        for k, sub in enumerate(SUBS):
            c0 = 3 + 5 * k
            sc, pr = rng(SC[sub]), PR[sub]
            if g:
                f = [f"=SUM({sc})", f"=AVERAGE({sc})",
                     f'=COUNTIF({sc},">="&参数设置!$C${pr})/$B{r}', f'=COUNTIF({sc},">="&参数设置!$D${pr})/$B{r}']
            else:
                f = [f"=SUMIF({cls},$A{r},{sc})", f"=AVERAGEIF({cls},$A{r},{sc})",
                     f'=COUNTIFS({cls},$A{r},{sc},">="&参数设置!$C${pr})/$B{r}',
                     f'=COUNTIFS({cls},$A{r},{sc},">="&参数设置!$D${pr})/$B{r}']
            f.append(score_formula(r, L(c0 + 1), L(c0 + 2), L(c0 + 3), pr))
            comps.append(f"{L(c0 + 4)}{r}")
            for j, (v, fmt) in enumerate(zip(f, ["0", "0.00", "0.0%", "0.0%", "0.00"])):
                c = s.cell(r, c0 + j, v)
                sty(c, j == 4, fl)
                c.number_format = fmt
        c = s.cell(r, cC, "=" + "+".join(comps)); sty(c, True, SUM); c.number_format = "0.00"
        sty(s.cell(r, cC + 1, "—" if g else f"=RANK({L(cC)}{r},${L(cC)}$4:${L(cC)}${3 + nc})"), True, SUM)
    note(s, 6 + nc, f"说明：及格线、优秀线、权重在「参数设置」中修改。当前为{cfg.scheme}方案，单科得分 = {cfg.formula_text()}。", ncol)
    s.column_dimensions["A"].width = 8; s.column_dimensions["B"].width = 8
    for c in range(3, cC):
        s.column_dimensions[L(c)].width = 8.5
    s.column_dimensions[L(cC)].width = 12; s.column_dimensions[L(cC + 1)].width = 6
    s.row_dimensions[3].height = 22
    s.freeze_panes = "C4"

    # ---------- 前85%分数段
    pp = f"{cfg.promote_ratio * 100:g}%"
    z = wb.create_sheet(f"前{pp}分数段", 1)
    title(z, f"{cfg.exam} 前{pp}学生分数段统计（按{NS}科总分，满分{cfg.total_full:g}）", 13)
    kf = [("实考人数", f"=COUNT({tot})"), (f"前{pp}对应名次", f"=ROUND(B3*{P['升学线比例（前百分之几）']},0)"),
          ("升学分数线（总分）", f"=LARGE({tot},B4)"), ("线上人数（含同分）", f'=COUNTIF({tot},">="&B5)'), ("线下人数", "=B3-B6")]
    for i, (k, v) in enumerate(kf):
        sty(z.cell(3 + i, 1, k), True, HEAD)
        sty(z.cell(3 + i, 2, v), True, SUM if i == 2 else None)
    z["C5"] = f"← 总分达到此分即进入全级前{pp}"
    z["C5"].font = Font(name=FONT, size=10, color="C00000")
    z["A9"] = f"一、线上学生分数段分布（每{cfg.band_width:g}分一段）"
    z["A9"].font = Font(name=FONT, size=12, bold=True)
    tc = 4 + nc
    for j, h in enumerate(["分数段", "下限", "上限(不含)"] + CL + ["全级人数", "占线上人数", "累计人数", "累计占全级"], 1):
        sty(z.cell(10, j, h), True, HEAD)
    r0 = 11
    fixed = R.bands[:-1]
    low_hi = R.bands[-1][2]
    for i in range(len(fixed) + 2):
        r = r0 + i
        last = i == len(fixed) + 1
        if i < len(fixed):
            lab, lo, hi = fixed[i][:3]
            sty(z.cell(r, 1, lab), True); sty(z.cell(r, 2, lo)); sty(z.cell(r, 3, hi))
        elif not last:
            sty(z.cell(r, 1, f'=B{r}&"–"&(C{r}-1)&"（线上最低段）"'), True)
            sty(z.cell(r, 2, "=$B$5")); sty(z.cell(r, 3, low_hi))
        else:
            sty(z.cell(r, 1, '="线下（"&$B$5&"分以下）"'), True, INPUT)
            sty(z.cell(r, 2, 0), fill=INPUT); sty(z.cell(r, 3, "=$B$5"), fill=INPUT)
        for j in range(nc):
            col = L(4 + j)
            f = (f'=COUNTIFS({cls},{col}$10,{tot},"<"&$B$5)' if last else
                 f'=COUNTIFS({cls},{col}$10,{tot},">="&MAX($B{r},$B$5),{tot},"<"&$C{r})')
            sty(z.cell(r, 4 + j, f), fill=INPUT if last else None)
        sty(z.cell(r, tc, f"=SUM(D{r}:{L(tc - 1)}{r})"), True, INPUT if last else None)
        if not last:
            c = z.cell(r, tc + 1, f"={L(tc)}{r}/$B$6"); sty(c); c.number_format = "0.0%"
            sty(z.cell(r, tc + 2, f"=SUM(${L(tc)}${r0}:{L(tc)}{r})"))
            c = z.cell(r, tc + 3, f"={L(tc + 2)}{r}/$B$3"); sty(c); c.number_format = "0.0%"
        else:
            c = z.cell(r, tc + 1, f"={L(tc)}{r}/$B$3"); sty(c, fill=INPUT); c.number_format = "0.0%"
            sty(z.cell(r, tc + 2, "占全级"), fill=INPUT); sty(z.cell(r, tc + 3, ""), fill=INPUT)
    rs = r0 + len(fixed) + 2
    sty(z.cell(rs, 1, "合计"), True, SUM)
    for j in range(2, tc + 4):
        sty(z.cell(rs, j, ""), fill=SUM)
    for j in range(4, tc + 1):
        sty(z.cell(rs, j, f"=SUM({L(j)}{r0}:{L(j)}{rs - 1})"), True, SUM)

    cuts = cfg.top_n
    def topn_block(start, ttl, pct_mode, ref=None):
        z.cell(start, 1, ttl).font = Font(name=FONT, size=12, bold=True)
        for j, h in enumerate(["名次段", "名次", "该名次总分"] + CL + ["实际人数（含同分）"], 1):
            sty(z.cell(start + 1, j, h), True, HEAD)
        for i, n in enumerate(cuts + [None]):
            r = start + 2 + i
            if n is None:
                sty(z.cell(r, 1, f"前{pp}（升学线）"), True, SUM); sty(z.cell(r, 2, "=$B$4"), fill=SUM)
            else:
                sty(z.cell(r, 1, f'="前"&B{r}&"名"'), True)
                if pct_mode:
                    sty(z.cell(r, 2, f"=B{ref + 2 + i}"))
                else:
                    sty(z.cell(r, 2, n), fill=INPUT, color="0000FF")
            sty(z.cell(r, 3, f"=LARGE({tot},B{r})"), True)
            for j in range(nc):
                col = L(4 + j)
                cnt = f'COUNTIFS({cls},{col}${start + 1},{tot},">="&$C{r})'
                if pct_mode:
                    c = z.cell(r, 4 + j, f"={cnt}/COUNTIF({cls},{col}${start + 1})"); sty(c); c.number_format = "0.0%"
                else:
                    sty(z.cell(r, 4 + j, "=" + cnt))
            sty(z.cell(r, 4 + nc, f'=COUNTIF({tot},">="&C{r})'), True)
    T1 = rs + 2
    topn_block(T1, "二、前N名各班人数（累计，含同分；黄色格子的名次可以改）", False)
    T2 = T1 + len(cuts) + 4
    topn_block(T2, "三、前N名占本班人数的比例（名次跟随上表）", True, T1)
    note(z, T2 + len(cuts) + 3, "说明：同分学生全部计入，所以实际人数可能比名次多。", 13)
    t2 = T2 + len(cuts) + 6
    z.cell(t2 - 1, 1, "四、各班上线情况").font = Font(name=FONT, size=12, bold=True)
    for j, h in enumerate(["班级", "实考人数", "线上人数", "上线率", "线下人数", "线上临界生", "线下临界生",
                           f"{cfg.high_line:g}分以上", "线上学生平均总分", "上线率名次"], 1):
        sty(z.cell(t2, j, h), True, HEAD)
    RG = P["临界生范围（分数线上下各几分）"]
    for i, cn in enumerate(CL + ["全级"]):
        r = t2 + 1 + i
        g = cn == "全级"
        fl = SUM if g else None
        cond = "" if g else f"{cls},$A{r},"
        sty(z.cell(r, 1, cn), True, fl)
        sty(z.cell(r, 2, f"=COUNT({tot})" if g else f"=COUNTIF({cls},$A{r})"), fill=fl)
        sty(z.cell(r, 3, f'=COUNTIFS({cond}{tot},">="&$B$5)'), True, fl)
        c = z.cell(r, 4, f"=C{r}/B{r}"); sty(c, True, fl); c.number_format = "0.0%"
        sty(z.cell(r, 5, f"=B{r}-C{r}"), fill=fl)
        sty(z.cell(r, 6, f'=COUNTIFS({cond}{tot},">="&$B$5,{tot},"<"&$B$5+{RG})'), fill=fl)
        sty(z.cell(r, 7, f'=COUNTIFS({cond}{tot},"<"&$B$5,{tot},">="&$B$5-{RG})'), fill=fl)
        sty(z.cell(r, 8, f'=COUNTIFS({cond}{tot},">={cfg.high_line:g}")'), fill=fl)
        c = z.cell(r, 9, f'=AVERAGEIFS({tot},{cond}{tot},">="&$B$5)'); sty(c, fill=fl); c.number_format = "0.0"
        sty(z.cell(r, 10, "—" if g else f"=RANK(D{r},$D${t2 + 1}:$D${t2 + nc})"), True, fl)
    note(z, t2 + nc + 3, "线上临界生：分数线到线上临界范围以内，要稳住；线下临界生：线下临界范围以内，是冲线重点。", 13)
    t3 = t2 + nc + 6
    z.cell(t3 - 1, 1, "五、线上 / 线下学生各科对比").font = Font(name=FONT, size=12, bold=True)
    for j, h in enumerate(["学科", "满分", "线上学生平均分", "线下学生平均分", "相差", "线上学生及格率", "线下学生及格率"], 1):
        sty(z.cell(t3, j, h), True, HEAD)
    for k, sub in enumerate(SUBS):
        r = t3 + 1 + k
        sc, pr = rng(SC[sub]), PR[sub]
        sty(z.cell(r, 1, sub), True); sty(z.cell(r, 2, f"=参数设置!$B${pr}"))
        for j, f, fm in [(3, f'=AVERAGEIF({tot},">="&$B$5,{sc})', "0.0"), (4, f'=AVERAGEIF({tot},"<"&$B$5,{sc})', "0.0"),
                         (5, f"=C{r}-D{r}", "0.0"),
                         (6, f'=COUNTIFS({tot},">="&$B$5,{sc},">="&参数设置!$C${pr})/$B$6', "0.0%"),
                         (7, f'=COUNTIFS({tot},"<"&$B$5,{sc},">="&参数设置!$C${pr})/$B$7', "0.0%")]:
            c = z.cell(r, j, f); sty(c); c.number_format = fm
    t4 = t3 + NS + 2
    z.cell(t4, 1, "六、分析结论（依据本次考试数据自动起草，请审核）").font = Font(name=FONT, size=12, bold=True)
    for i, line in enumerate(conclusions):
        c = z.cell(t4 + 1 + i, 1, f"{i + 1}. {line}")
        c.font = Font(name=FONT, size=11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        z.merge_cells(start_row=t4 + 1 + i, start_column=1, end_row=t4 + 1 + i, end_column=13)
        z.row_dimensions[t4 + 1 + i].height = 36
    z.column_dimensions["A"].width = 22
    for col in "BCDEFGHIJKLM":
        z.column_dimensions[col].width = 10
    z.column_dimensions["I"].width = 14

    # ---------- 临界生名单
    q = wb.create_sheet("临界生名单", 2)
    title(q, f"临界生名单（升学线上下{cfg.near_range:g}分以内，按总分从高到低）", 8 + NS, 14)
    note(q, 2, "名单按本次成绩生成；「距分数线」「最薄弱学科」随分数线自动计算。最薄弱学科按得分率（成绩 ÷ 满分）比较。", 8 + NS)
    for j, h in enumerate(["班级", "姓名", "考号"] + SUBS + ["总分", "级名次", "位置", "距分数线", "最薄弱学科"], 1):
        sty(q.cell(3, j, h), True, HEAD)
    hc0 = 10 + NS
    for k, sub in enumerate(SUBS):
        q.cell(3, hc0 + k, sub).font = Font(name=FONT, size=9, color="999999")
    q.cell(2, hc0, "得分率（辅助列）").font = Font(name=FONT, size=9, color="999999")
    cut_ref = f"'前{pp}分数段'!$B$5"
    TT = L(4 + NS)
    for i, row in enumerate(R.near.itertuples(index=False), start=4):
        vals = [row.班级, row.姓名, row.考号] + [getattr(row, sub) for sub in SUBS]
        for j, v in enumerate(vals, 1):
            sty(q.cell(i, j, v))
        sty(q.cell(i, 4 + NS, f"=SUM(D{i}:{L(3 + NS)}{i})"), True)
        sty(q.cell(i, 5 + NS, f"=SUMPRODUCT(--({tot}>{TT}{i}))+1"))
        sty(q.cell(i, 6 + NS, f'=IF({TT}{i}>={cut_ref},"线上","线下")'), True)
        sty(q.cell(i, 7 + NS, f"={TT}{i}-{cut_ref}"))
        for k, sub in enumerate(SUBS):
            c = q.cell(i, hc0 + k, f"={L(4 + k)}{i}/参数设置!$B${PR[sub]}")
            c.number_format = "0%"
            c.font = Font(name=FONT, size=9, color="999999")
        hr = f"${L(hc0)}{i}:${L(hc0 + NS - 1)}{i}"
        sty(q.cell(i, 8 + NS, f"=INDEX(${L(hc0)}$3:${L(hc0 + NS - 1)}$3,MATCH(MIN({hr}),{hr},0))"), True, color="C00000")
    for j in range(1, 9 + NS):
        q.column_dimensions[L(j)].width = 8.5
    q.column_dimensions[L(8 + NS)].width = 11
    q.freeze_panes = "D4"

    # ---------- 任课安排 + 教师排名
    if cfg.teachers:
        MAXC = max(len(t.classes) for t in cfg.teachers)
        t = wb.create_sheet("任课安排")
        title(t, "任课安排（黄色格子可修改，教师排名会自动更新）", 2 + MAXC, 14)
        note(t, 2, f"班级写法与成绩明细一致（如 {cfg.grade_char}1、{cfg.grade_char}2）。", 2 + MAXC)
        for j, h in enumerate(["学科", "教师"] + [f"任教班级{i + 1}" for i in range(MAXC)], 1):
            sty(t.cell(3, j, h), True, HEAD)
        TROW = {}
        for i, tc_ in enumerate(cfg.teachers):
            r = 4 + i
            TROW.setdefault(tc_.subject, []).append(r)
            sty(t.cell(r, 1, tc_.subject), True)
            sty(t.cell(r, 2, tc_.name), fill=INPUT, color="0000FF")
            for j in range(MAXC):
                sty(t.cell(r, 3 + j, tc_.classes[j] if j < len(tc_.classes) else None), fill=INPUT, color="0000FF")
        for j in range(1, 3 + MAXC):
            t.column_dimensions[L(j)].width = 11
        e = wb.create_sheet("教师排名", 1)
        title(e, f"{cfg.exam} 任课教师成绩排名", 10)
        note(e, 2, "按所教班级的学生合并计算（算法与各班统计相同）；名次为同学科教师之间的排名。", 10)
        row = 4
        for sub in SUBS:
            rows_t = TROW.get(sub)
            if not rows_t:
                continue
            sc, pr = rng(SC[sub]), PR[sub]
            for j, h in enumerate(["学科", "教师", "任教班级", "实考人数", "总分", "平均分", "及格率", "优秀率", "综合得分", "学科内名次"], 1):
                sty(e.cell(row, j, h), True, HEAD)
            first, lastr = row + 1, row + len(rows_t)
            for i, ar in enumerate(rows_t):
                r = row + 1 + i
                cc = [f"任课安排!${L(3 + j)}${ar}" for j in range(MAXC)]
                agg = lambda fn: "+".join(f'IF({c}="",0,{fn(c)})' for c in cc)
                sty(e.cell(r, 1, sub), True)
                sty(e.cell(r, 2, f"=任课安排!$B${ar}"), True)
                sty(e.cell(r, 3, "=" + cc[0] + "".join(f'&IF({c}="","","、"&{c})' for c in cc[1:])))
                sty(e.cell(r, 4, "=" + agg(lambda c: f"COUNTIF({cls},{c})")))
                c = e.cell(r, 5, "=" + agg(lambda c: f"SUMIF({cls},{c},{sc})")); sty(c); c.number_format = "0"
                c = e.cell(r, 6, f"=E{r}/D{r}"); sty(c); c.number_format = "0.00"
                c = e.cell(r, 7, "=(" + agg(lambda c: f'COUNTIFS({cls},{c},{sc},">="&参数设置!$C${pr})') + f")/D{r}")
                sty(c); c.number_format = "0.0%"
                c = e.cell(r, 8, "=(" + agg(lambda c: f'COUNTIFS({cls},{c},{sc},">="&参数设置!$D${pr})') + f")/D{r}")
                sty(c); c.number_format = "0.0%"
                c = e.cell(r, 9, score_formula(r, "F", "G", "H", pr)); sty(c, True, SUM); c.number_format = "0.00"
                sty(e.cell(r, 10, f"=RANK(I{r},$I${first}:$I${lastr})"), True, SUM)
            row = lastr + 2
        for col, w in zip("ABCDEFGHIJ", [7, 10, 20, 9, 9, 9, 9, 9, 10, 11]):
            e.column_dimensions[col].width = w
    sl = None
    if W:      # 结构分单独一张表，不放进「各班统计」（用户明确要求两张表分开）
        sl = _structure_sheet(wb, R, P, rng, cls, tot, SC, NS, TC, rT)
    if compare:
        _compare_sheets(wb, R, compare, {"rng": rng, "cls": cls, "tot": tot, "name": rng("C"), "rank": rng(L(6 + NS)),
                                         "PR": PR, "cC": cC, "t2": t2, "pp": pp, "struct": sl})
    wb.save(path)
    return path


def _structure_sheet(wb, R: Result, P, rng, cls, tot, SC, NS, TC, rT):
    """班级结构分（带公式）。"""
    cfg, CL, W = R.cfg, R.classes, R.cfg.structure
    nc = len(CL)
    allp, alle = rng(L(8 + NS)), rng(L(9 + NS))
    cut_ref = f"'前{cfg.promote_ratio * 100:g}%分数段'!$B$5"
    TF = f"参数设置!$B${rT}"
    w = lambda k: P[f"结构分：{k}" + ("（分值）" if k in RATE_ITEMS or k == "增值评价" else "")]

    # 列定义：(组名, 小标题, 公式模板（{r} = 行号；{X} = 本行某列的列字母）, 格式, 是否黄色可改)
    cols = [("", "班级", None, None, False)]
    if "参考率" in W:
        cols.append(("", "应考人数", "enrolled", "0", True))
    cols.append(("", "实考人数", "=COUNTIF(%s,$A{r})" % cls, "0", False))
    pts = []
    for k in RATE_ITEMS:
        if k not in W:
            continue
        g = f"{k}（{W[k]:g}分）"
        if k == "平均成绩":
            cols += [(g, "总分平均", "=AVERAGEIF(%s,$A{r},%s)" % (cls, tot), "0.00", False),
                     (g, "得分率", "={总分平均}{r}/" + TF, "0.0%", False)]
            rate = "得分率"
        elif k in ("全科合格率", "全科优秀率"):
            lab = k[:4] + "人数"
            cols += [(g, lab, "=COUNTIFS(%s,$A{r},%s,1)" % (cls, allp if k == "全科合格率" else alle), "0", False),
                     (g, "比率", "={%s}{r}/{实考人数}{r}" % lab, "0.0%", False)]
            rate = "比率"
        elif k == "参考率":
            cols += [(g, "参考率", "={实考人数}{r}/{应考人数}{r}", "0.0%", False)]
            rate = "参考率"
        else:  # 进线率
            cols += [(g, "进线人数", '=COUNTIFS(%s,$A{r},%s,">="&%s)' % (cls, tot, cut_ref), "0", False),
                     (g, "比率", "={进线人数}{r}/{实考人数}{r}", "0.0%", False)]
            rate = "比率"
        key = f"{k}分"
        cols.append((g, key, "={%s}{r}*%s" % (f"{g}|{rate}", w(k)), "0.00", False))
        pts.append(key)
    cols.append(("", "小计", "=" + "+".join("{%s}{r}" % k for k in pts), "0.00", False))
    tail = ["小计"]
    if "增值评价" in W:
        g = f"增值评价（{W['增值评价']:g}分）"
        base = "+".join("{%s分}{r}" % k for k in TWO_RATES if k in W) or "0"
        cols += [(g, "两率一分", "=" + base, "0.00", False),
                 (g, "两率一分名次", "=RANK({两率一分}{r},{两率一分}$5:{两率一分}$%d)" % (4 + nc), "0", False),
                 (g, "上次名次", "prev", "0", True),
                 (g, "增值评价分", "=MIN(%s,%s-%s*({两率一分名次}{r}-1)+IF({上次名次}{r}=\"\",0,({上次名次}{r}-{两率一分名次}{r})*%s))"
                  % (w("增值评价"), w("增值评价"), P["结构分：增值名次差"], P["结构分：增值进退步"]), "0.00", False)]
        tail.append("增值评价分")
    if "前10名每人加分" in W:
        g = f"前10名加分（每人{W['前10名每人加分']:g}分）"
        cols += [(g, "前10名人数", '=COUNTIFS(%s,$A{r},%s,">="&LARGE(%s,MIN(10,COUNT(%s))))' % (cls, tot, tot, tot), "0", False),
                 (g, "加分", "={前10名人数}{r}*%s" % P["结构分：前10名每人加分"], "0.00", False)]
        tail.append("加分")
    cols += [("", "结构分", "=" + "+".join("{%s}{r}" % k for k in tail), "0.00", False),
             ("", "名次", "=RANK({结构分}{r},{结构分}$5:{结构分}$%d)" % (4 + nc), "0", False)]

    # 列字母索引：小标题唯一的直接用小标题，重名的（比率）用“组名|小标题”
    letter = {}
    for j, (g, h, *_) in enumerate(cols, 1):
        letter[f"{g}|{h}"] = L(j)
        letter.setdefault(h, L(j))

    def fill(tpl, r):
        out = tpl
        for k in sorted(letter, key=len, reverse=True):
            out = out.replace("{%s}" % k, letter[k])
        return out.replace("{r}", str(r))

    ws = wb.create_sheet("班级结构分", 1)
    ncol = len(cols)
    title(ws, f"{cfg.exam} 班级结构分评比（{cfg.scheme}）", ncol)
    note(ws, 2, "黄色格子可以修改（应考人数、上次名次），分值在「参数设置」中修改，其他自动重新计算。", ncol)
    j = 1
    while j <= ncol:
        g = cols[j - 1][0]
        k = j
        while k < ncol and cols[k][0] == g and g:
            k += 1
        if g:
            sty(ws.cell(3, j, g), True, HEAD)
            ws.merge_cells(start_row=3, start_column=j, end_row=3, end_column=k)
            for jj in range(j, k + 1):
                sty(ws.cell(4, jj, cols[jj - 1][1]), True, HEAD)
        else:
            fl = SUM if cols[j - 1][1] in ("结构分", "名次", "小计") else HEAD
            sty(ws.cell(3, j, cols[j - 1][1]), True, fl)
            ws.merge_cells(start_row=3, start_column=j, end_row=4, end_column=j)
            ws.cell(4, j).border = BORDER
        j = k + 1
    for i, c in enumerate(CL):
        r = 5 + i
        x = R.structure[c]
        for jj, (g, h, tpl, fmt, inp) in enumerate(cols, 1):
            if tpl is None:
                v = c
            elif tpl == "enrolled":                     # 没填应考人数：这一格先等于实考人数，填上实际人数即可重算
                v = fill("={实考人数}{r}", r) if x.get("应考按实考") else x["应考人数"]
            elif tpl == "prev":
                v = x["上次名次"]
            else:
                v = fill(tpl, r)
            cell = ws.cell(r, jj, v)
            sty(cell, h in ("班级", "结构分", "名次"), INPUT if inp else (SUM if h in ("结构分", "名次") else None),
                "0000FF" if inp else "000000")
            if fmt:
                cell.number_format = fmt
    lines = [f"说明：各项直接按比例折分（比率 × 分值）。平均成绩 = 本班总分平均分 ÷ 总分满分（{cfg.total_full:g}）× 分值；"
             "全科合格 / 全科优秀 = 每一科都达到合格线 / 优秀线的学生。"]
    if "参考率" in W:
        lines.append("应考人数 = 应该参加考试的在籍学生；实考人数 = 各科成绩齐全、计入统计的学生"
                     "（缺考、请假、成绩不全的不算）。参考率 = 实考人数 ÷ 应考人数。"
                     + ("没有填应考人数的班，应考人数按实考人数算（参考率 100%）；在黄色格子里填上实际人数就会重算。"
                        if any(R.structure[c].get("应考按实考") for c in CL) else ""))
    if "进线率" in W:
        lines.append(f"进线率 = 本班升学线（全级前{cfg.promote_ratio:.0%}）以上人数 ÷ 实考人数。")
    if "增值评价" in W:
        lines.append(f"增值评价：按“两率一分”（平均成绩、全科合格率、全科优秀率三项得分之和，不含参考率）排名，第1名得 {W['增值评价']:g} 分，"
                     f"每低一名减 {W.get('增值名次差', 0.1):g} 分；再与上次考试的两率一分名次比，每进步一名加 {W.get('增值进退步', 0.05):g} 分、"
                     f"每退步一名减 {W.get('增值进退步', 0.05):g} 分，最高 {W['增值评价']:g} 分。上次名次空着 = 不算进步退步。"
                     "下次考试时，把本表的“两率一分名次”作为上次名次输入。")
    if "前10名每人加分" in W:
        lines.append(f"前10名加分：全级总分前10名（同分计入）每有一人，本班加 {W['前10名每人加分']:g} 分。")
    for i, t in enumerate(lines):
        note(ws, 7 + nc + i, t, ncol)
        ws.row_dimensions[7 + nc + i].height = 30
    ws.column_dimensions["A"].width = 7
    for jj in range(2, ncol + 1):
        ws.column_dimensions[L(jj)].width = 9.5
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "B5"
    return letter


UP_DOWN = '"↑"0;"↓"0;"持平"'          # 名次进退：正数是进步
DIFF2, DIFFP = "+0.00;-0.00;0.00", "+0.0%;-0.0%;0.0%"


def _heads(ws, r, heads, widths=None):
    for j, h in enumerate(heads, 1):
        sty(ws.cell(r, j, h), True, HEAD)
    for j, w in enumerate(widths or [], 1):
        ws.column_dimensions[L(j)].width = w


def _put(ws, r, j, v, fmt=None, bold=False, fill=None, color="000000"):
    c = ws.cell(r, j, v)
    sty(c, bold, fill, color)
    if fmt:
        c.number_format = fmt
    return c


def _compare_sheets(wb, R: Result, C: dict, ref: dict):
    """和上次比：上次的数字写成固定值（当时算好的结果），本次的数字用公式取自本工作簿的其他表。"""
    cfg, CL = R.cfg, R.classes
    nc, pv = len(CL), C["prev"]
    GREY = "808080"
    pp = ref["pp"]
    cut_ref = f"'前{pp}分数段'!$B$5"
    match = lambda sheet, col, r0, r1, key: (f"INDEX({sheet}!${col}${r0}:${col}${r1},"
                                             f"MATCH({key},{sheet}!$A${r0}:$A${r1},0))")
    who = f"上次：{pv['exam']}" + (f"（{pv['date']}）" if pv["date"] else "") + f"，{pv['scheme']}方案，实考 {pv['n']} 人，升学线 {pv['cut']:g} 分"

    # ---------- 和上次比：各班、各科、上次临界生汇总
    ws = wb.create_sheet("和上次比")
    title(ws, f"{cfg.exam} 和上次考试比较", 13)
    note(ws, 2, who + "。灰色是上次的结果（固定数字），黑色是本次（公式，随本表其他工作表变化）。“进退”正数为进步。", 13)
    ws.row_dimensions[2].height = 30
    ws.cell(4, 1, "一、各班").font = Font(name=FONT, size=12, bold=True)
    heads = ["班级"]
    if C["hasStructure"]:
        heads += ["上次结构分", "本次结构分", "变化", "上次结构分名次", "本次结构分名次", "进退"]
    heads += ["上次综合名次", "本次综合名次", "进退", "上次上线率", "本次上线率", "变化"]
    _heads(ws, 5, heads, [16] + [11] * 13)
    ws.row_dimensions[5].height = 30
    sl = ref["struct"]
    for i, x in enumerate(C["classes"]):
        r, j = 6 + i, 1
        _put(ws, r, 1, x["name"], bold=True)
        if C["hasStructure"]:
            cur = match("班级结构分", sl["|结构分"], 5, 4 + nc, f"$A{r}")
            curk = match("班级结构分", sl["|名次"], 5, 4 + nc, f"$A{r}")
            f = x["found"] and "prevStruct" in x
            _put(ws, r, 2, x["prevStruct"] if f else "—", "0.00", color=GREY)
            _put(ws, r, 3, "=" + cur, "0.00", True)
            _put(ws, r, 4, f"=C{r}-B{r}" if f else "", DIFF2)
            _put(ws, r, 5, x["prevStructRank"] if f else "—", "0", color=GREY)
            _put(ws, r, 6, "=" + curk, "0", True)
            _put(ws, r, 7, f"=E{r}-F{r}" if f else "", UP_DOWN)
            j = 7
        f = x["found"]
        cur = match("各班统计", L(ref["cC"] + 1), 4, 3 + nc, f"$A{r}")
        _put(ws, r, j + 1, x["prevRank"] if f else "—", "0", color=GREY)
        _put(ws, r, j + 2, "=" + cur, "0", True)
        _put(ws, r, j + 3, f"={L(j + 1)}{r}-{L(j + 2)}{r}" if f else "", UP_DOWN)
        cur = match(f"'前{pp}分数段'", "D", ref["t2"] + 1, ref["t2"] + nc, f"$A{r}")
        _put(ws, r, j + 4, x["prevOnlineRate"] if f else "—", "0.0%", color=GREY)
        _put(ws, r, j + 5, "=" + cur, "0.0%", True)
        _put(ws, r, j + 6, f"={L(j + 5)}{r}-{L(j + 4)}{r}" if f else "上次没有这个班", DIFFP)
    r = 6 + nc + 1
    ws.cell(r, 1, "二、各科（全年级）").font = Font(name=FONT, size=12, bold=True)
    note(ws, r + 1, "两次考试的满分可能不同，所以平均分比“得分率”（平均分 ÷ 满分）。“变化”是百分点。", 13)
    h0 = r + 2
    _heads(ws, h0, ["科目", "上次满分", "本次满分", "上次平均分", "本次平均分", "上次得分率", "本次得分率", "变化",
                    "上次及格率", "本次及格率", "变化", "上次优秀率", "本次优秀率", "变化"])
    ws.row_dimensions[h0].height = 30
    grow = 4 + nc                                            # 各班统计里“全年级”那一行
    for k, x in enumerate(C["subjects"]):
        r = h0 + 1 + k
        c0 = 3 + 5 * cfg.subjects.index(x["subject"])
        _put(ws, r, 1, x["subject"], bold=True)
        _put(ws, r, 2, x["prevFull"], "0", color=GREY)
        _put(ws, r, 3, f"=参数设置!$B${ref['PR'][x['subject']]}", "0")
        _put(ws, r, 4, x["prevAvg"], "0.00", color=GREY)
        _put(ws, r, 5, f"=各班统计!{L(c0 + 1)}{grow}", "0.00")
        _put(ws, r, 6, f"=D{r}/B{r}", "0.0%", color=GREY)
        _put(ws, r, 7, f"=E{r}/C{r}", "0.0%", True)
        _put(ws, r, 8, f"=G{r}-F{r}", DIFFP)
        _put(ws, r, 9, x["prevPass"], "0.0%", color=GREY)
        _put(ws, r, 10, f"=各班统计!{L(c0 + 2)}{grow}", "0.0%", True)
        _put(ws, r, 11, f"=J{r}-I{r}", DIFFP)
        _put(ws, r, 12, x["prevExc"], "0.0%", color=GREY)
        _put(ws, r, 13, f"=各班统计!{L(c0 + 3)}{grow}", "0.0%", True)
        _put(ws, r, 14, f"=M{r}-L{r}", DIFFP)
    if not C["subjects"]:
        note(ws, h0 + 1, "两次考试没有相同的科目。", 13)
    r = h0 + 2 + max(len(C["subjects"]), 1)
    ws.cell(r, 1, "三、上次的临界生，这次怎么样了").font = Font(name=FONT, size=12, bold=True)
    _heads(ws, r + 1, ["上次", "人数", "这次能对上", "其中", "人数"])
    end = 3 + max(len(C["near"]), 1)
    nC, nG = f"上次临界生!$C$4:$C${end}", f"上次临界生!$G$4:$G${end}"
    for i, (lab, pos, goal) in enumerate([("上次线下临界生", "线下", "上线了"), ("上次线上临界生", "线上", "掉到线下")]):
        rr = r + 2 + i
        t = C["nearUp"] if pos == "线下" else C["nearDown"]
        _put(ws, rr, 1, lab, bold=True)
        _put(ws, rr, 2, t["all"], "0", color=GREY)
        _put(ws, rr, 3, f'=COUNTIF({nC},"{pos}")-COUNTIFS({nC},"{pos}",{nG},"这次没有对上")', "0")
        _put(ws, rr, 4, goal)
        _put(ws, rr, 5, f'=COUNTIFS({nC},"{pos}",{nG},"{goal}")', "0", True)
    note(ws, r + 5, "学生按“班级 + 姓名”对应：两次都参加、并且能对上的才比较；换了班、同班重名、缺考的不在比较之内。"
                    "逐人的情况见「名次进退」「上次临界生」两张表。", 13)
    ws.row_dimensions[r + 5].height = 30
    ws.freeze_panes = "B6"

    # ---------- 名次进退：每个能对上的学生
    m = wb.create_sheet("名次进退")
    title(m, f"{cfg.exam} 学生级名次进退（和 {pv['exam']} 比）", 7, 14)
    note(m, 2, "按进步名次从多到少排列，可以用表头的筛选按钮按班级筛选或重新排序。本次总分、名次取自「成绩明细」。", 7)
    _heads(m, 3, ["班级", "姓名", "上次总分", "本次总分", "上次级名次", "本次级名次", "进退（名）"], [8, 10, 10, 10, 11, 11, 11])
    look = lambda col, r: f"=SUMIFS({col},{ref['cls']},$A{r},{ref['name']},$B{r})"
    for i, x in enumerate(C["students"]):
        r = 4 + i
        _put(m, r, 1, x["cls"]); _put(m, r, 2, x["name"], bold=True)
        _put(m, r, 3, x["prevTotal"], None, color=GREY)
        _put(m, r, 4, look(ref["tot"], r))
        _put(m, r, 5, x["prevRank"], "0", color=GREY)
        _put(m, r, 6, look(ref["rank"], r), "0", True)
        _put(m, r, 7, f"=E{r}-F{r}", UP_DOWN, True)
    m.freeze_panes = "C4"
    m.auto_filter.ref = f"A3:G{3 + max(len(C['students']), 1)}"

    # ---------- 上次临界生：逐人的去向
    q = wb.create_sheet("上次临界生")
    title(q, f"上次（{pv['exam']}）的临界生，这次怎么样了", 7, 14)
    note(q, 2, f"本次升学线见「前{pp}分数段」（{R.cut:g} 分），改了升学线比例，“本次位置”“去向”会跟着变。名单含学生姓名，仅限教师使用。", 7)
    _heads(q, 3, ["班级", "姓名", "上次位置", "上次总分", "本次总分", "本次位置", "去向"], [8, 10, 10, 10, 10, 10, 13])
    for i, x in enumerate(C["near"]):
        r = 4 + i
        _put(q, r, 1, x["cls"]); _put(q, r, 2, x["name"], bold=True)
        _put(q, r, 3, x["prevPos"], color=GREY); _put(q, r, 4, x["prevTotal"], None, color=GREY)
        if x["found"]:
            _put(q, r, 5, look(ref["tot"], r))
            _put(q, r, 6, f'=IF(E{r}>={cut_ref},"线上","线下")', bold=True)
            _put(q, r, 7, f'=IF(C{r}="线下",IF(F{r}="线上","上线了","仍在线下"),IF(F{r}="线下","掉到线下","仍在线上"))', bold=True)
        else:
            _put(q, r, 5, ""); _put(q, r, 6, "")
            _put(q, r, 7, "这次没有对上", color=GREY)
    q.freeze_panes = "C4"
