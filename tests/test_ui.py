"""界面的端到端测试：用自动化浏览器像真人一样操作（导入 → 设置 → 计算 → 查看 → 导出 → 试用到期 → 激活）。
需要 playwright（没装就跳过）。只用 samples/ 里的虚构数据，在临时文件夹里进行，不碰真实的 data/ 和 output/。"""
import datetime as dt
import shutil
import sys
from pathlib import Path

import pytest

pw = pytest.importorskip("playwright.sync_api")
ROOT = Path(__file__).resolve().parent.parent


@pytest.fixture()
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("CHENGJI_NO_OPEN", "1")
    (tmp_path / "下载").mkdir()
    monkeypatch.setenv("CHENGJI_DOWNLOADS", str(tmp_path / "下载"))          # 模板不要真的存到用户的“下载”里
    for d in ("config", "templates", "samples"):
        shutil.copytree(ROOT / d, tmp_path / d)
    (tmp_path / "data").mkdir()
    shutil.copy2(ROOT / "samples" / "示例任课总表.xlsx", tmp_path / "data" / "任课总表.xlsx")
    (tmp_path / "data" / "本校设置.yaml").write_text("已引导: true\n", encoding="utf-8")
    from chengji.server import App, start
    httpd, url, a = start(App(tmp_path))
    yield tmp_path, url
    httpd.shutdown()


@pytest.fixture()
def page():
    with pw.sync_playwright() as p:
        try:
            b = p.chromium.launch()
        except Exception as e:                      # 没装浏览器内核
            pytest.skip(f"没有可用的浏览器内核：{e}")
        pg = b.new_page(viewport={"width": 1320, "height": 880})
        errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        yield pg
        b.close()
        assert not errs, errs


def settings(pg, tab="school"):
    """打开设置页的某个标签：school 学校与科目、algo 分数线与算法、teacher 任课总表、look 外观、lic 使用许可。"""
    pg.evaluate("go('settings')"); pg.click(f'#stabs [data-s="{tab}"]')
    pg.wait_for_selector(f'.spane[data-s="{tab}"].on')


def idle(pg):
    pg.wait_for_function("!document.querySelector('#busy').classList.contains('on')", timeout=120000)


def import_and_compute(pg, url, sample, exam, fill=None):
    pg.goto(url)
    pg.wait_for_selector("#recent .li")
    pg.set_input_files("#file", str(ROOT / "samples" / sample))
    pg.wait_for_selector("#wizBody:not(.hide)")
    pg.fill("#exam", exam)
    if fill:
        fill(pg)
    pg.click("#run")


def test_full_flow(app, page):
    root, url = app
    def fill(pg):
        for c, v in zip(["九1", "九2", "九3", "九4", "九5", "九6"], [3, 6, 1, 5, 4, 2]):
            pg.fill(f'#prev input[data-c="{c}"]', str(v))
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "界面测试", fill)
    page.wait_for_selector("#resBody:not(.hide)")
    assert page.inner_text("#rTitle") == "界面测试"
    first = page.inner_text("#ovRank .rk").split()
    assert first[:3] == ["1", "九6", "46.45"]                      # 与命令行、Excel 公式核对过的结果一致
    assert "进步 1 名" in page.inner_text("#ovRank .rk")
    page.click('#tabs button[data-t="de"]')
    page.click('#fc button:has-text("九2")')
    assert page.inner_text("#cnt") == "共 45 人"
    assert page.inner_text("#deT tbody tr td") == "1"              # 筛选后序号从 1 排起
    page.click('#tabs button[data-t="te"]')
    assert page.locator("#teG .tcard").count() == 7
    page.evaluate("go('export')")
    page.click("#xlsBtn"); idle(page)
    out = root / "output" / "界面测试"
    assert (out / "界面测试_各班综合统计.xlsx").stat().st_size > 20000
    assert (out / "结果.json").is_file() and (out / "登分表副本.xlsx").is_file()
    # 重新打开程序后，能从“最近的考试”里回看并导出
    page.goto(url); page.wait_for_selector('#recent .li[data-n="界面测试"]')
    page.click('#recent .li[data-n="界面测试"]'); page.wait_for_selector("#resBody:not(.hide)")
    assert page.inner_text("#ovRank .rk").split()[:3] == ["1", "九6", "46.45"]


def test_zhongkao_needs_full_marks(app, page):
    """中考 10 科：生物、地理、体育没定满分 → 先拦住并说明；填上后算出进线率、前 10 名加分。"""
    root, url = app
    def fill(pg):
        pg.click('#scheme button:has-text("中考核算")')
    import_and_compute(page, url, "示例登分表_中考.xlsx", "中考测试", fill)
    page.wait_for_selector("#wizErr:not(.hide)")
    assert "生物、地理、体育" in page.inner_text("#wizErr") and "满分" in page.inner_text("#wizErr")
    for s in ("生物", "地理", "体育"):
        page.fill(f'#fullChips input[data-s="{s}"]', "50")
    page.click("#run"); page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="st"]')
    head = page.inner_text("#stT thead")
    assert "进线率" in head and "前 10 名加分" in head and "增值评价" not in head


def test_settings_saved_locally(app, page):
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    settings(page); page.wait_for_selector("#setChips .chip")
    page.fill("#sSchool", "测试中学"); page.fill('#setChips input[data-s="生物"]', "50")
    page.click("#saveBtn"); page.wait_for_selector("#saved.on")
    text = (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")
    assert "测试中学" in text and "生物: 50" in text
    assert "××中学" in (root / "config" / "学校设置.yaml").read_text(encoding="utf-8")   # 公共设置文件不被改动


def test_trial_expired_then_activate(app, page, tmp_path, monkeypatch):
    """试用到期：能算能看、不能导出；输入本机的激活码后可以导出。（需要本机有许可模块和生成器）"""
    impl = pytest.importorskip("chengji._license_impl")
    if not (ROOT / "admin" / "工具" / "make_code.py").is_file():
        pytest.skip("没有发码程序")
    sys.path.insert(0, str(ROOT / "admin" / "工具"))
    import make_code
    monkeypatch.setenv("CHENGJI_LICENSE_ENFORCE", "1")
    monkeypatch.setenv("CHENGJI_LICENSE_DIR", str(tmp_path / "lic"))
    monkeypatch.setenv("CHENGJI_MACHINE_RAW", "ui-test-machine")
    monkeypatch.setattr(impl, "_MACHINE", None)
    impl.status()                                                   # 开始试用
    for d in impl._dirs():                                          # 把“第一次使用”挪到 4 天前
        st = impl._unseal((d / "state.json").read_text()); st["first"] -= 4 * 86400
        (d / "state.json").write_text(impl._seal(st))
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "到期测试")
    page.wait_for_selector("#resBody:not(.hide)")                   # 到期后仍然能算、能看
    assert page.inner_text("#licText") == "试用已结束"
    page.evaluate("go('export')")
    assert page.is_visible("#export .lock")                         # 导出被锁住
    assert not (root / "output" / "到期测试" / "到期测试_各班综合统计.xlsx").exists()
    assert "微信" in page.inner_text("#export .lock .contactBox")        # 锁定提示里写着找谁要激活码
    page.click("#export .lock button.pri"); page.wait_for_selector("#mask.on")
    assert page.inner_text("#machine") == impl.machine_code()
    assert "申请激活码请联系" in page.inner_text("#mask .contactBox")
    page.fill("#code", "AAAAA-BBBBB"); page.click("#actBtn")
    page.wait_for_function("document.querySelector('#actErr').textContent.length > 0")
    code = make_code.make(impl.machine_code(), dt.date.today() + dt.timedelta(days=365))
    page.fill("#code", code); page.click("#actBtn")
    page.wait_for_function("!document.querySelector('#mask').classList.contains('on')")
    assert page.inner_text("#licText") == "已激活" and not page.is_visible("#export .lock")
    page.click("#xlsBtn"); idle(page)
    assert (root / "output" / "到期测试" / "到期测试_各班综合统计.xlsx").is_file()


def test_recent_list_is_short(app, page):
    """首页“最近的考试”只显示最近 5 次，更多的点“查看全部”再展开。"""
    import json
    root, url = app
    for i in range(8):
        d = root / "output" / f"考试{i}"
        d.mkdir(parents=True)
        (d / "结果.json").write_text(json.dumps({"savedAt": f"2026-10-0{i + 1} 08:00", "opts": {},
                                               "view": {"exam": f"考试{i}", "grade": "九年级", "scheme": "平时", "n": 100, "date": ""}}, ensure_ascii=False), encoding="utf-8")
    page.goto(url); page.wait_for_selector("#moreRecent")
    assert page.locator("#recent .li[data-n]").count() == 5
    assert "查看全部 8 次" in page.inner_text("#moreRecent")
    page.click("#moreRecent"); page.wait_for_function("document.querySelectorAll('#recent .li[data-n]').length === 8")
    page.click("#moreRecent"); page.wait_for_function("document.querySelectorAll('#recent .li[data-n]').length === 5")


def test_enrolment_can_be_left_blank(app, page):
    """没有任课总表、不填应考人数也能算：按实考人数，结果里有说明。"""
    root, url = app
    (root / "data" / "任课总表.xlsx").unlink()
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "不填应考")
    page.wait_for_selector("#resBody:not(.hide)")
    assert "没有填应考人数，按实考人数计算" in page.inner_text("#ovNotes")
    page.click('#tabs button[data-t="st"]')
    assert "按实考人数计算" in page.inner_text("#stNote")
    assert page.inner_text("#stT tbody tr").split()[2:4] == ["44", "44"]      # 九6：应考 = 实考 = 44


def test_template_download(app, page, tmp_path):
    """主页面点模板。浏览器方式下走浏览器自己的下载；内容就是那张空白模板。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#tpls a")
    assert page.locator("#tpls a").all_inner_texts() == ["七年级", "八年级", "九年级", "中考"]
    with page.expect_download() as d:
        page.click('#tpls a[data-k="九年级"]')
    assert d.value.suggested_filename == "登分表_九年级.xlsx"
    d.value.save_as(tmp_path / "a.xlsx")
    assert (tmp_path / "a.xlsx").read_bytes() == (root / "templates" / "登分表_九年级.xlsx").read_bytes()
    with page.expect_download() as d:
        page.click("#tplT")
    assert d.value.suggested_filename == "任课总表.xlsx"
    assert not page.is_visible("#wizBody")                           # 点模板不会触发导入


def test_template_save_dialog(tmp_path, monkeypatch):
    """程序窗口里：弹“存储”对话框，存到使用者选的位置；没写扩展名自动补上；点取消什么也不做。"""
    from chengji.server import App
    for d in ("config", "templates"):
        shutil.copytree(ROOT / d, tmp_path / d)
    app = App(tmp_path)
    asked = []

    class FakeWindow:                                                # 冒充程序窗口：记下对话框的参数，返回“用户选的路径”
        def __init__(self, answer):
            self.answer = answer

        def create_file_dialog(self, kind, directory="", save_filename="", file_types=()):
            asked.append((directory, save_filename))
            return self.answer

    app.window = FakeWindow((str(tmp_path / "我的桌面" / "九年级期中登分"),))
    (tmp_path / "我的桌面").mkdir()
    r = app.api_template({"key": "九年级"})
    assert asked[-1][1] == "登分表_九年级.xlsx"                       # 对话框里预填的文件名
    saved = tmp_path / "我的桌面" / "九年级期中登分.xlsx"
    assert r["name"] == saved.name and saved.read_bytes() == (tmp_path / "templates" / "登分表_九年级.xlsx").read_bytes()
    app.window = FakeWindow(None)                                    # 点了取消
    assert app.api_template({"key": "任课总表"}) == {"cancelled": True}
    assert asked[-1][1] == "任课总表.xlsx" and len(list((tmp_path / "我的桌面").iterdir())) == 1


def test_new_subject_in_wizard(app, page, tmp_path):
    """登分表里加了一门新科目：界面上标“新”、要求填满分；可以点掉不计；算过一次后设置里记住它。"""
    import openpyxl
    root, url = app
    wb = openpyxl.load_workbook(ROOT / "samples" / "示例登分表_九年级.xlsx"); ws = wb.active
    c = ws.max_column + 1
    ws.cell(2, c, "信息技术"); ws.cell(2, c + 1, "加试")
    for r in range(3, ws.max_row + 1):
        ws.cell(r, c, 30 + r % 20); ws.cell(r, c + 1, r % 10)
    f = tmp_path / "加科目.xlsx"; wb.save(f)
    page.goto(url); page.wait_for_selector("#recent .li")
    page.set_input_files("#file", str(f)); page.wait_for_selector("#wizBody:not(.hide)")
    assert "9 科" in page.inner_text("#insKv") and "新加的科目" in page.inner_text("#insNotes")
    assert page.locator('#fullChips .chip.need').count() == 2
    page.fill("#exam", "加科目测试")
    page.click('#fullChips .chip[data-s="加试"] .x')                 # “加试”这次不算
    page.click("#run"); page.wait_for_selector("#wizErr:not(.hide)")
    assert "信息技术 还没有填满分" in page.inner_text("#wizErr") and "加试" not in page.inner_text("#wizErr")
    page.fill('#fullChips input[data-s="信息技术"]', "50")
    page.click("#run"); page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="cl"]')
    head = page.inner_text("#clT thead")
    assert "信息技术" in head and "加试" not in head
    assert "信息技术: 50" in (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")     # 记住了，下次自动认得
    settings(page); page.wait_for_selector('#setChips input[data-s="信息技术"]')
    page.fill("#newSub", "劳动"); page.fill("#newSubFull", "40"); page.click("#saveBtn"); page.wait_for_selector('#setChips input[data-s="劳动"]')
    assert "劳动: 40" in (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")


def test_algo_editable_in_settings(app, page):
    """设置里改算法：合计不是 100% 不让存；存了之后新算的考试按新比例，旧考试不变；可以恢复默认。"""
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "改比例之前")
    page.wait_for_selector("#resBody:not(.hide)")
    before = page.inner_text("#ovRank .rk").split()[:3]
    settings(page, "algo"); page.wait_for_selector("#algoStruct .chip")
    assert page.inner_text("#sum1") == "合计 100%" and page.inner_text("#sum2") == "合计 100%"
    assert [i.input_value() for i in page.locator("#algoStruct input").all()] == ["50", "20", "20", "5", "0", "5"]
    page.fill('#algoStruct input[data-k="平均成绩"]', "60")
    assert "合计 110%" in page.inner_text("#sum2")
    assert page.is_visible("#savebar")                                                    # 改了就出现“还没有保存”
    page.click("#saveBtn"); page.wait_for_function("document.querySelector('#saveErr').textContent.includes('110%')")
    assert "算法方案" not in (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")   # 没存进去
    page.fill('#algoStruct input[data-k="全科优秀率"]', "10")                             # 60 + 20 + 10 + 5 + 0 + 5 = 100
    page.fill('#algoScore input[data-k="优秀率"]', "30"); page.fill('#algoScore input[data-k="及格率"]', "20")
    page.click("#saveBtn"); page.wait_for_selector("#saved.on")
    text = (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")
    assert "平均成绩: 60" in text and "优秀率: 0.3" in text
    assert page.is_visible("#algoTag") and not page.is_visible("#savebar")
    # 新算一次：按新比例
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "改比例之后")
    page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="st"]')
    assert "平均成绩 / 60" in page.inner_text("#stT thead") and "全科优秀率 / 10" in page.inner_text("#stT thead")
    # 旧考试回看、导出：仍是旧比例
    page.goto(url); page.wait_for_selector('#recent .li[data-n="改比例之前"]')
    page.click('#recent .li[data-n="改比例之前"]'); page.wait_for_selector("#resBody:not(.hide)")
    assert page.inner_text("#ovRank .rk").split()[:3] == before
    page.evaluate("go('export')"); page.click("#xlsBtn"); idle(page)
    import openpyxl
    ws = openpyxl.load_workbook(root / "output" / "改比例之前" / "改比例之前_各班综合统计.xlsx")["班级结构分"]
    assert any("平均成绩（50分）" in str(c.value) for c in ws[3])
    # 恢复默认
    settings(page, "algo"); page.wait_for_selector("#algoStruct .chip")
    page.click("#algoReset"); page.wait_for_selector("#dlg.on"); page.click("#dlgOk")
    page.wait_for_function("document.querySelector('#algoStruct input').value === '50'")
    assert "算法方案" not in (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")


def test_settings_state_the_rules(app, page):
    """设置界面要把两条规则写明白：平时折算百分制、中考不折算；增值评价封顶。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    settings(page, "algo"); page.wait_for_selector("#algoStruct .chip")
    f = page.inner_text("#algoFormula")                                    # 平时考试
    assert page.is_checked("#algoNorm") and "优秀率×25% + 及格率×25% + 平均分÷满分×100×50%" in f and "折算成百分制" in f
    add = page.inner_text("#algoAdd")
    assert "第 1 名得满分 5 分" in add and "最高不超过 5 分" in add and "超过的也按 5 分算" in add
    page.click('#algoSeg button[data-n="中考"]')                             # 中考核算
    f = page.inner_text("#algoFormula")
    assert not page.is_checked("#algoNorm") and "及格率×40% + 平均分（原始分，不折算）×60%" in f and "优秀率" not in f.split("\n")[0]
    assert not page.is_visible("#algoAdd")                                 # 中考默认没有增值评价
    page.fill('#algoStruct input[data-k="增值评价"]', "8")                   # 改成 8% → 封顶跟着变成 8 分
    assert "最高不超过 8 分" in page.inner_text("#algoAdd")


def test_welcome_on_first_run(app, page):
    """第一次打开：出现欢迎页，填学校名称后进入；之后不再出现。"""
    root, url = app
    (root / "data" / "本校设置.yaml").unlink()                         # 全新的电脑
    page.goto(url); page.wait_for_selector("#welcome.on")
    page.fill("#wSchool", "某某县第一中学"); page.click("#wOk")
    page.wait_for_function("document.querySelector('#schoolName').textContent === '某某县第一中学'")
    text = (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")
    assert "学校: 某某县第一中学" in text and "落款: 某某县第一中学{年级}组" in text
    page.goto(url); page.wait_for_selector("#recent .li"); page.wait_for_timeout(300)
    assert not page.is_visible("#welcome")


def test_try_sample(app, page):
    """“用示例数据试一试”：一点就出结果，标明是虚构数据，不进“最近的考试”，教师排名用的是示例教师。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#tryBtn")
    page.click("#tryBtn"); page.wait_for_selector("#resBody:not(.hide)")
    assert page.is_visible("#sampleTag") and page.inner_text("#ovRank .rk").split()[:3] == ["1", "九6", "46.45"]
    page.click('#tabs button[data-t="te"]')
    assert "示例教师" in page.inner_text("#teG")
    page.evaluate("go('home')"); page.wait_for_selector("#recent .li")
    assert page.locator("#recent .li[data-n]").count() == 0


def test_sort_by_header(app, page):
    """成绩明细点表头排序：分数先高到低，再点低到高，第三次恢复；序号始终从 1 排起。"""
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "排序测试")
    page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="de"]')
    total = lambda: [float(x) for x in page.locator("#deT tbody tr td:nth-child(12)").all_inner_texts()]
    orig = total()
    page.click('#deT th:has-text("总分")')
    assert total() == sorted(orig, reverse=True) and page.inner_text("#deT tbody tr td") == "1"
    assert page.inner_text("#deT tbody tr td:nth-child(13)") == "1"          # 排第一行的就是级名次第 1
    page.click('#deT th:has-text("总分")'); assert total() == sorted(orig)
    page.click('#deT th:has-text("总分")'); assert total() == orig
    page.click('#fc button:has-text("九2")'); page.click('#deT th:has-text("班名次")')
    ranks = [int(x) for x in page.locator("#deT tbody tr td:nth-child(14)").all_inner_texts()]
    assert ranks == sorted(ranks) and ranks[0] == 1 and len(ranks) == 45


def test_prefill_from_last_exam(app, page):
    """同年级再算一次：上次的结构分名次自动带进来，可以清空。"""
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "第一次")
    page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="st"]')
    first = {r.split()[1]: r.split()[0] for r in page.locator("#stT tbody tr").all_inner_texts()}     # 班级 -> 名次
    page.goto(url); page.wait_for_selector("#recent .li")
    page.set_input_files("#file", str(ROOT / "samples" / "示例登分表_九年级.xlsx")); page.wait_for_selector("#wizBody:not(.hide)")
    assert "已带入《第一次》" in page.inner_text("#prevHint")
    assert {c: page.input_value(f'#prev input[data-c="{c}"]') for c in first} == first
    page.click("#prevClear")
    assert all(i.input_value() == "" for i in page.locator("#prev input").all())


def test_rename_and_delete_exam(app, page):
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "要改名的")
    page.wait_for_selector("#resBody:not(.hide)")
    page.evaluate("go('export')"); page.click("#xlsBtn"); idle(page)
    assert page.is_visible("#xlsDone") and "要改名的_各班综合统计.xlsx" in page.inner_text("#xlsDone")   # 导出后有“打开文件 / 打开文件夹”
    page.evaluate("go('home')"); page.wait_for_selector('#recent .li[data-n="要改名的"]')
    page.hover('#recent .li[data-n="要改名的"]'); page.click('#recent .li[data-n="要改名的"] .ren')
    page.wait_for_selector("#dlg.on"); page.fill("#dlgI", "改好了"); page.click("#dlgOk")
    page.wait_for_selector('#recent .li[data-n="改好了"]')
    assert sorted(f.name for f in (root / "output" / "改好了").iterdir()) == ["改好了_各班综合统计.xlsx", "登分表副本.xlsx", "结果.json"]
    assert not (root / "output" / "要改名的").exists()
    page.hover('#recent .li[data-n="改好了"]'); page.click('#recent .li[data-n="改好了"] .del')
    page.wait_for_selector("#dlg.on"); page.click("#dlgOk")
    page.wait_for_function("document.querySelectorAll('#recent .li[data-n]').length === 0")
    assert not (root / "output" / "改好了").exists()
    assert len(list((root / "output" / "已删除").iterdir())) == 1               # 没有真删，移到了“已删除”


def test_unsaved_settings_warning_and_font(app, page):
    """设置改了没保存就离开：提醒；字号调大后记住。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    settings(page); page.wait_for_selector("#setChips .chip")
    assert not page.is_visible("#savebar")
    page.fill("#sSchool", "临时改的名字")
    assert page.is_visible("#savebar")
    page.click('#nav button[data-p="home"]'); page.wait_for_selector("#dlg.on")
    page.click("#dlgAlt")                                                      # 不保存
    page.wait_for_function("document.querySelector('#home').classList.contains('on')")
    assert "临时改的名字" not in (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")
    settings(page, "school"); page.fill("#sSchool", "又改了一次")
    assert page.is_visible('#stabs [data-s="school"] .dot')                   # 标签上的小金点：这里有没保存的修改
    page.click("#discardBtn"); page.wait_for_selector('#stabs [data-s="school"] .dot', state="hidden")


def test_appearance(app, page):
    """外观：主题（跟随系统 / 明亮 / 暗色）、字号放大缩小、恢复字号、恢复默认外观。立即生效，记在本机设置里，不走保存条。"""
    root, url = app
    yaml_text = lambda: (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")
    theme = lambda: page.evaluate("document.documentElement.dataset.theme")
    zoom = lambda: page.evaluate("getComputedStyle(document.documentElement).getPropertyValue('--zoom').trim()")
    page.emulate_media(color_scheme="light")
    page.goto(url); page.wait_for_selector("#recent .li")
    assert theme() == "light" and zoom() == "1"
    page.emulate_media(color_scheme="dark"); page.wait_for_timeout(100)
    assert theme() == "dark"                                                   # 跟随系统：系统变深色，界面跟着变
    settings(page, "look")
    assert page.inner_text("#zoomVal") == "100%" and page.is_disabled("#zoomReset") and page.is_disabled("#lookReset")
    page.click('#themeSeg button[data-v="明亮"]'); assert theme() == "light"
    page.click("#zoomIn"); page.click("#zoomIn"); page.wait_for_timeout(500)
    assert page.inner_text("#zoomVal") == "120%" and zoom() == "1.2"
    assert "界面主题: 明亮" in yaml_text() and "界面字号: 120" in yaml_text()
    assert not page.is_visible("#savebar")                                     # 外观立即生效，不算“未保存的修改”
    page.click("#zoomReset"); page.wait_for_timeout(500)
    assert zoom() == "1" and "界面字号" not in yaml_text() and "界面主题: 明亮" in yaml_text()
    page.keyboard.press("Control+Equal"); page.keyboard.press("Control+Equal"); page.keyboard.press("Control+Minus")
    page.wait_for_timeout(500)
    assert page.inner_text("#zoomVal") == "110%" and "界面字号: 110" in yaml_text()
    page.click('#themeSeg button[data-v="暗色"]'); page.wait_for_timeout(500)
    page.emulate_media(color_scheme="light")
    page.goto(url); page.wait_for_selector("#recent .li")                     # 重新打开：一开始就是暗色、110%（服务直接填进页面）
    assert theme() == "dark" and zoom() == "1.1"
    settings(page, "look"); assert "暗色" in page.inner_text("#lookNow")
    page.click("#lookReset"); page.wait_for_timeout(400)
    assert theme() == "light" and zoom() == "1" and "界面主题" not in yaml_text() and "界面字号" not in yaml_text()
    assert page.is_disabled("#lookReset")


def test_tabs_are_clear_and_keyboard_friendly(app, page):
    """标签：选中的有 aria-selected，下面一行说明；左右方向键能切换；带 ? 号的表头点一下出说明（不排序）。"""
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "标签测试")
    page.wait_for_selector("#resBody:not(.hide)")
    assert page.get_attribute('#tabs [data-t="ov"]', "aria-selected") == "true" and "主要结论" in page.inner_text("#tabHint")
    page.click('#tabs [data-t="st"]')
    assert "满分 100" in page.inner_text("#tabHint") and page.is_visible("#st") and not page.is_visible("#ov")
    page.keyboard.press("ArrowRight")
    assert page.get_attribute('#tabs [data-t="cl"]', "aria-selected") == "true" and page.is_visible("#cl")
    assert page.get_attribute('#tabs [data-t="te"]', "data-tip")             # 没选中的标签停一下也有说明
    page.click('#tabs [data-t="de"]')
    order = page.inner_text("#deT tbody tr td:nth-child(3)")
    page.click('#deT th:has-text("级名次") .qi')
    page.wait_for_selector("#tip.on"); assert "全年级" in page.inner_text("#tip")
    assert page.inner_text("#deT tbody tr td:nth-child(3)") == order          # 点 ? 号只出说明，不排序
    settings(page, "algo")
    assert "比例" in page.inner_text("#stabHint") and page.get_attribute('#stabs [data-s="algo"]', "aria-selected") == "true"
    settings(page, "teacher"); page.click("#tEdit"); page.wait_for_selector("#ted.on")
    page.keyboard.press("Escape"); page.wait_for_selector("#ted", state="hidden")  # Esc 关掉弹窗


def _second_exam(tmp_path):
    """造“第二次考试”：在九年级示例的基础上，九1 全班每科加一些分、九6 减一些分。"""
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / "samples" / "示例登分表_九年级.xlsx"); ws = wb.active
    for r in range(3, ws.max_row + 1):
        cls = str(ws.cell(r, 2).value)
        for c in range(5, 12):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                full = [120, 100, 100, 70, 60, 50, 50][c - 5]
                if cls == "九1": ws.cell(r, c, min(full, round(v * 1.15) + 2))
                if cls == "九6": ws.cell(r, c, max(0, round(v * 0.85)))
    f = tmp_path / "第二次.xlsx"; wb.save(f)
    return f


def test_class_detail_and_compare(app, page, tmp_path):
    """点班级看详情；同年级算过两次后，“和上次比”里能看到各班、各科、学生的变化。"""
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "第一次月考")
    page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="cp"]')
    assert page.is_visible("#cpEmpty")                                   # 只算过一次：没得比
    page.click('#tabs button[data-t="ov"]'); page.click('#ovRank .clk[data-c="九6"]')
    page.wait_for_selector("#drawer.on")
    body = page.inner_text("#drawerBody")
    assert "九6" in body and "全级第 1 名" in body and "各科情况" in body and "本班前 10 名" in body and "临界生" in body
    assert page.locator("#drawerBody table").first.locator("tbody tr").count() == 7       # 7 科
    page.click("#toDetail"); page.wait_for_function("!document.querySelector('#drawer').classList.contains('on')")
    assert page.inner_text("#cnt") == "共 44 人"                           # 跳到成绩明细并筛出这个班
    # 第二次考试
    page.goto(url); page.wait_for_selector("#recent .li")
    page.set_input_files("#file", str(_second_exam(tmp_path))); page.wait_for_selector("#wizBody:not(.hide)")
    page.fill("#exam", "期中考试"); page.click("#run"); page.wait_for_selector("#resBody:not(.hide)")
    assert "进步" in page.inner_text('#ovRank')                           # 上次名次自动带入 → 增值评价里有进退步
    page.click('#tabs button[data-t="cp"]'); page.wait_for_selector("#cpBody:not(.hide)")
    page.wait_for_function("document.querySelector('#cpCls tbody tr') !== null")
    assert "第一次月考" in page.inner_text("#cpSel")
    cls = {r.split()[0]: r for r in page.locator("#cpCls tbody tr").all_inner_texts()}
    assert "↑" in cls["九1"] and "↓" in cls["九6"]                        # 九1 进步、九6 退步
    assert page.locator("#cpSub tbody tr").count() == 7
    up, dn = page.inner_text("#cpUp"), page.inner_text("#cpDn")
    assert "九1" in up and "九6" in dn and "↑" in up and "↓" in dn
    assert "这次有" in page.inner_text("#cpNear") and "能对上的共" in page.inner_text("#cpNote")


def test_teacher_editor(app, page):
    """在软件里直接改任课总表：加一位教师、改应考人数，保存后程序读到的就是新的。"""
    from chengji.config import load_enrolled, load_school, load_teachers
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    settings(page, "teacher"); page.wait_for_selector("#tEdit")
    page.click("#tEdit"); page.wait_for_selector("#ted.on")
    page.click('#tedSeg button[data-g="九年级"]')
    n = page.locator("#tedRows tr[data-i]").count()
    page.click("#tedAdd")
    row = page.locator("#tedRows tr[data-i]").last
    row.locator('input[data-f="subject"]').fill("语文"); row.locator('input[data-f="teacher"]').fill("新来的老师")
    page.click("#tedOk"); page.wait_for_function("document.querySelector('#tedErr').textContent.includes('一起填')")     # 班级没填：拦住并说明
    row = page.locator("#tedRows tr[data-i]").last
    row.locator('input[data-f="classes"]').fill("九1 九3")
    page.locator('#tedCls .chip input[data-f="enrolled"]').first.fill("52")
    page.click("#tedOk"); page.wait_for_function("!document.querySelector('#ted').classList.contains('on')")
    school = load_school(root / "config" / "学校设置.yaml")
    T = load_teachers(root / "data" / "任课总表.xlsx", school, "九年级")
    assert len(T) == n + 1 and ("语文", "新来的老师", ["九1", "九3"]) in [(t.subject, t.name, t.classes) for t in T]
    assert load_enrolled(root / "data" / "任课总表.xlsx", school, "九年级")["九1"] == 52
    assert (root / "data" / "任课总表_上一份.xlsx").is_file()             # 改之前的那份留着
    assert "九年级 31 位" in page.inner_text("#tInfo")


def test_help_and_lines(app, page):
    """帮助页能打开；分数线在设置里能改，合计、高低关系有检查；鼠标停在表头上有解释。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    assert page.title() == "分寸 · 成绩核算" and page.inner_text("#brandName") == "分寸"
    page.click('#nav button[data-p="help"]')
    txt = page.inner_text("#help")
    for k in ("三步上手", "登分表怎么填", "名词解释", "全科合格率", "常见问题", "关于"):
        assert k in txt
    assert page.inner_text("#aboutVer") and not page.is_visible("#h5")       # 源码运行没有激活模块：不显示“试用与激活”
    settings(page, "algo"); page.wait_for_selector("#lineChips .chip")
    assert [i.input_value() for i in page.locator("#lineChips input").all()] == ["60", "80", "85", "75", "15"]
    page.fill('#lineChips input[data-l="及格线"]', "85"); page.click("#saveBtn")
    page.wait_for_function("document.querySelector('#saveErr').textContent.includes('优秀线要比及格线高')")
    page.fill('#lineChips input[data-l="及格线"]', "50"); page.fill('#lineChips input[data-l="临界范围"]', "10")
    page.click("#saveBtn"); page.wait_for_selector("#saved.on")
    text = (root / "data" / "本校设置.yaml").read_text(encoding="utf-8")
    assert "及格线比例: 0.5" in text and "临界范围: 10" in text
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "新分数线")
    page.wait_for_selector("#resBody:not(.hide)")
    page.click('#tabs button[data-t="st"]')
    assert "每一科都达到合格线" in page.get_attribute('#stT th:has-text("全科合格率")', "data-tip")
    page.click('#tabs button[data-t="ne"]')
    assert all(abs(int(v)) <= 10 for v in page.locator("#neD tbody tr td:nth-child(4), #neU tbody tr td:nth-child(4)").all_inner_texts())


def test_wizard_steps_all_visible(app, page):
    """向导顶部三步都要看得见（曾经因为样式名相撞，第一步被隐藏了）。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    page.set_input_files("#file", str(ROOT / "samples" / "示例登分表_九年级.xlsx")); page.wait_for_selector("#wizBody:not(.hide)")
    steps = page.locator(".steps .step")
    assert steps.count() == 3 and all(steps.nth(i).is_visible() for i in range(3))
    assert ["导入登分表" in steps.nth(0).inner_text(), "确认设置" in steps.nth(1).inner_text(), "查看结果" in steps.nth(2).inner_text()] == [True] * 3


def test_contact_author_in_help(app, page):
    """帮助页有“联系作者”：微信号和项目的 GitHub 网址，可以复制。激活窗口保持原样，不放二维码。"""
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li")
    page.click('#nav button[data-p="help"]')
    assert "联系作者" in page.locator(".toc").inner_text()
    page.click('.toc a[data-h="h8"]')
    assert page.inner_text("#cWechat") == "915274394"
    assert page.inner_text("#cGithub") == "github.com/lumanman996/chengji-tool"
    page.click("#cWechatCopy"); page.wait_for_selector("#toast.on")
    assert "已复制微信号" in page.inner_text("#toast")
    assert page.locator("#help img").count() == 0 and page.locator("#mask img").count() == 0      # 没有二维码


def test_update_prompt(app, page, monkeypatch):
    """打开程序时查到新版本：弹出提示，点“立即更新”就下载、换上新版本、退出（之后由替换脚本重新打开）。"""
    import os
    import threading
    from chengji import updater
    root, url = app
    monkeypatch.setenv("CHENGJI_UPDATE_CHECK", "1")
    info = {"ok": True, "current": "2.0.2", "latest": "9.9.9", "newer": True, "notes": "· 新功能一\n· 新功能二", "page": "",
            "canAuto": True, "reason": "", "error": "", "asset": {"name": "chengji-tool-v9.9.9-mac.zip"}}
    monkeypatch.setattr(updater, "check", lambda *a, **k: dict(info))
    calls = []
    monkeypatch.setattr(updater.Updater, "start", lambda self, asset: (calls.append(("start", asset["name"])), self._set(phase="ready", percent=100)))
    monkeypatch.setattr(updater.Updater, "apply", lambda self, args=None: calls.append(("apply",)))
    exited = threading.Event()
    monkeypatch.setattr(os, "_exit", lambda code: exited.set())                 # 真程序会在这里退出
    page.goto(url)
    page.wait_for_selector("#upd.on", timeout=10000)
    assert "9.9.9" in page.inner_text("#updP") and "新功能一" in page.inner_text("#updNotes")
    page.click("#updGo")
    page.wait_for_function("document.querySelector('#updTip').textContent.includes('重新打开')")
    assert exited.wait(5) and calls == [("start", "chengji-tool-v9.9.9-mac.zip"), ("apply",)]
    # 不能一键更新时（比如没有写入权限）：说明原因，只给“打开下载页”
    info.update(canAuto=False, reason="程序所在的文件夹不允许修改")
    page.goto(url); page.wait_for_selector("#upd.on", timeout=10000)
    assert not page.is_visible("#updGo") and "不允许修改" in page.inner_text("#updErr")
    page.click("#updLater"); assert not page.is_visible("#upd")
    # 帮助页手动检查：已经是最新
    info.update(newer=False, latest="2.0.2")
    page.evaluate("go('help')"); page.click("#updBtn")
    page.wait_for_function("document.querySelector('#updMsg').textContent.includes('已经是最新')")


def test_no_update_check_by_default(app, page, monkeypatch):
    """源码运行（开发、测试）时打开程序不联网检查。"""
    from chengji import updater
    monkeypatch.setattr(updater, "check", lambda *a, **k: pytest.fail("不该检查"))
    root, url = app
    page.goto(url); page.wait_for_selector("#recent .li"); page.wait_for_timeout(1800)
    assert not page.is_visible("#upd")


def test_export_with_compare(app, page, tmp_path):
    """导出页选“和上次比”：Excel 多三张对比表，PDF 里可以勾“和上次比”；选“不加对比”就和原来一样。"""
    import openpyxl
    root, url = app
    import_and_compute(page, url, "示例登分表_九年级.xlsx", "第一次月考")
    page.wait_for_selector("#resBody:not(.hide)")
    page.goto(url); page.wait_for_selector("#recent .li")
    page.set_input_files("#file", str(_second_exam(tmp_path))); page.wait_for_selector("#wizBody:not(.hide)")
    page.fill("#exam", "期中考试"); page.click("#run"); page.wait_for_selector("#resBody:not(.hide)")
    page.evaluate("go('export')")
    assert page.is_visible("#expCmpRow") and page.input_value("#expCmp") == "第一次月考"
    ck = page.locator('#pdfList input[value="和上次比"]')
    assert ck.count() == 1 and not ck.is_checked()                  # 含学生姓名，默认不勾
    page.click("#xlsBtn"); idle(page)
    x = root / "output" / "期中考试" / "期中考试_各班综合统计.xlsx"
    assert {"和上次比", "名次进退", "上次临界生"} <= set(openpyxl.load_workbook(x).sheetnames)
    page.click('#pdfList label:has(input[value="和上次比"])'); assert ck.is_checked()
    page.click("#pdfBtn"); idle(page)
    if "已导出" in page.inner_text("#pdfDone"):                       # 电脑上没有浏览器时不生成 PDF
        assert (root / "output" / "期中考试" / "期中考试_成绩发布版.pdf").stat().st_size > 20000
    page.select_option("#expCmp", "")
    assert ck.is_disabled() and not ck.is_checked()
    page.click("#xlsBtn"); idle(page)
    assert "和上次比" not in openpyxl.load_workbook(x).sheetnames
