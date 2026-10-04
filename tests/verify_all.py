"""一键全面核对：用虚构示例数据把各年级、各方案都算一遍，
用 LibreOffice 重算导出的 Excel，再逐项对照程序的独立计算结果。必须 0 处不一致。
最后两种情况带“和上次比”：先算一次“上次”，再算“本次”并导出对比表，一起核对。

用法：python tests/verify_all.py        （需要先装好 LibreOffice）
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))
from verify_excel import verify, verify_compare   # noqa: E402

T = str(ROOT / "samples" / "示例任课总表.xlsx")
PREV = "九1=3 九2=6 九3=1 九4=5 九5=4 九6=2"
# (登分表, 方案, 升学比例, 上次名次, 满分)
FULL = "生物=50 地理=50 体育=50"      # 这三科在学校设置里满分还没定，示例数据按 50 分造的
CASES = [("七年级", "平时", None, "", FULL),
         ("八年级", "平时", None, "", FULL),
         ("九年级", "平时", None, PREV, ""),
         ("九年级", "中考", 0.8, "", ""),          # 只有 7 科的表用中考方案（会提示少了 3 科，照样要算对）
         ("中考", "中考", 0.8, "", FULL),
         ("九年级+新科目", "平时", None, PREV, "信息技术=50"),   # 登分表里加一门科目总表没有的科目
         ("九年级+改比例", "平时", None, PREV, ""),               # 在界面里改过比例（平时方案加了进线率、前10名加分）
         ("九年级+改比例", "中考", 0.8, PREV, "")]                # 中考方案加了增值评价、去掉参考率


# 和上次比：(本次的方案, 升学比例, 本次满分, 本次加不加一门新科目)。“上次”都是九年级示例、平时方案。
COMPARE_CASES = [("平时", None, "", False),
                 ("中考", 0.8, "数学=150 信息技术=50", True)]       # 方案、满分、科目都和上次不一样


def second_exam(dst: Path, new_subject: bool) -> str:
    """造“本次考试”：九年级示例每人每科上下浮动几分；再故意造几种对不上的情况：
    改一个人的名字、让一个班里出现两个同名、让一个人缺考、把一个人换到别的班。"""
    import openpyxl
    wb = openpyxl.load_workbook(ROOT / "samples" / "示例登分表_九年级.xlsx")
    ws = wb.active
    full = [120, 100, 100, 70, 60, 50, 50]
    for r in range(3, ws.max_row + 1):
        for k, c in enumerate(range(5, 12)):
            v = ws.cell(r, c).value
            if isinstance(v, (int, float)):
                ws.cell(r, c, max(0, min(full[k], v + (r * 7 + c * 13) % 15 - 6)))
    ws.cell(3, 3, str(ws.cell(3, 3).value) + "新")                   # 改名：对不上
    ws.cell(6, 3, ws.cell(5, 3).value)                               # 同班重名：两人都不比
    ws.cell(8, 5, "缺考")                                            # 缺考：不计入，也不比
    ws.cell(10, 2, "九2")                                            # 换班：对不上
    if new_subject:
        c = ws.max_column + 1
        ws.cell(2, c, "信息技术")
        for r in range(3, ws.max_row + 1):
            ws.cell(r, c, 20 + (r * 7) % 31)
    wb.save(dst)
    return str(dst)


def compare_cases(soffice: str, tmp: Path) -> tuple[int, int]:
    from chengji import service as sv
    from chengji.config import load_school
    total = bad = 0
    school = load_school(ROOT / "config" / "学校设置.yaml")
    out = tmp / "cmp"
    prev_file = ROOT / "samples" / "示例登分表_九年级.xlsx"
    ranks = dict(x.split("=") for x in PREV.split())
    for i, (scheme, ratio, full, new_subject) in enumerate(COMPARE_CASES):
        R0, c0, m0 = sv.compute(prev_file, school, {"scheme": "平时", "exam": f"上次{i}", "prev": ranks}, T)
        sv.save_run(out, prev_file, {"scheme": "平时"}, sv.to_view(R0, c0, m0))
        cur = second_exam(tmp / f"本次{i}.xlsx", new_subject)
        fulls = dict(x.split("=") for x in full.split())
        R1, c1, _m = sv.compute(cur, school, {"scheme": scheme, "exam": f"本次{i}", "ratio": ratio, "full": fulls, "prev": ranks}, T)
        C = sv.compare_with(R1, out, f"上次{i}")
        xlsx = sv.export_excel(R1, c1, out, C)
        subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "re" / f"cmp{i}"), str(xlsx)],
                       check=True, capture_output=True, timeout=300)
        re_ = tmp / "re" / f"cmp{i}" / xlsx.name
        n1, e1 = verify(cur, re_, T, scheme, ratio, "", PREV, full)
        n2, e2 = verify_compare(re_, R1, C)
        total += n1 + n2
        bad += len(e1) + len(e2)
        print(f"和上次比（本次{scheme}方案{'，改了满分、加了新科目' if new_subject else ''}）：核对 {n1 + n2} 项"
              f"（其中对比表 {n2} 项，能对上的学生 {len(C['students'])} 人），不一致 {len(e1) + len(e2)} 项")
        for e in (e1 + e2)[:20]:
            print("   ", e)
    return total, bad


def find_soffice() -> str:
    for c in ("soffice", "libreoffice", "/Applications/LibreOffice.app/Contents/MacOS/soffice",
              r"C:\Program Files\LibreOffice\program\soffice.exe"):
        w = shutil.which(c) or (c if Path(c).is_file() else None)
        if w:
            return w
    sys.exit("没有找到 LibreOffice（soffice），无法重算 Excel 公式。")


def main():
    soffice = find_soffice()
    total = bad = 0
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for i, (grade, scheme, ratio, prev, full) in enumerate(CASES):
            scores = str(ROOT / "samples" / f"示例登分表_{grade}.xlsx")
            if grade.endswith("+新科目"):                       # 临时造一张：九年级示例后面加一列“信息技术”
                import openpyxl
                wb = openpyxl.load_workbook(ROOT / "samples" / "示例登分表_九年级.xlsx")
                ws = wb.active
                c = ws.max_column + 1
                ws.cell(2, c, "信息技术")
                for r in range(3, ws.max_row + 1):
                    ws.cell(r, c, 20 + (r * 7) % 31)
                scores = str(tmp / "加了新科目.xlsx")
                wb.save(scores)
            name = f"核对{i}"
            cmd = [sys.executable, "-m", "chengji", scores, "--任课", T, "--考试", name, "--方案", scheme,
                   "--输出", str(tmp / "out"), "--不生成PDF", "--不确认"]
            cmd += ["--升学比例", str(ratio)] if ratio else []
            cmd += ["--上次名次", prev] if prev else []
            cmd += ["--满分", full] if full else []
            local = None
            if grade.endswith("+改比例"):                       # 界面里改过比例：单科得分 30/20/50，结构分各项都用上
                import yaml
                local = tmp / "本校设置.yaml"
                local.write_text(yaml.safe_dump({
                    "算法方案": {"平时": {"优秀率": 0.3, "及格率": 0.2, "平均分": 0.5, "平均分折算百分制": "是"},
                                 "中考": {"优秀率": 0.1, "及格率": 0.3, "平均分": 0.6, "平均分折算百分制": "否"}},
                    "结构分方案": {"平时": {"平均成绩": 40, "全科合格率": 22.5, "全科优秀率": 17.5, "参考率": 5, "进线率": 10, "增值评价": 5,
                                            "增值名次差": 0.2, "增值进退步": 0.1, "前10名每人加分": 0.3},
                                   "中考": {"平均成绩": 60, "全科合格率": 20, "全科优秀率": 10, "增值评价": 10, "增值名次差": 0.5, "增值进退步": 0.25}}},
                    allow_unicode=True), encoding="utf-8")
                scores = str(ROOT / "samples" / "示例登分表_九年级.xlsx")
                cmd[3] = scores
                cmd += ["--本校设置", str(local)]
            subprocess.run(cmd, check=True, cwd=ROOT, capture_output=True)
            xlsx = tmp / "out" / name / f"{name}_各班综合统计.xlsx"
            subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(tmp / "re" / name), str(xlsx)],
                           check=True, capture_output=True, timeout=300)
            n, errs = verify(scores, tmp / "re" / name / xlsx.name, T, scheme, ratio, "", prev, full, local=local)
            total += n
            bad += len(errs)
            print(f"{grade} {scheme}方案：核对 {n} 项，不一致 {len(errs)} 项")
            for e in errs[:20]:
                print("   ", e)
        n, b = compare_cases(soffice, tmp)
        total += n
        bad += b
    print(f"合计核对 {total} 项，不一致 {bad} 项")
    sys.exit(1 if bad else 0)


if __name__ == "__main__":
    main()
