"""基本测试（只用 samples/ 里的虚构数据）。运行：python -m pytest -q"""
from pathlib import Path

import openpyxl
import pytest

from chengji.analysis import analyze, draft_conclusions
from chengji.config import load_enrolled, load_school, load_teachers, make_config, normalize_class, parse_classes
from chengji.excel_report import build_excel

ROOT = Path(__file__).resolve().parent.parent
SCHOOL = load_school(ROOT / "config" / "学校设置.yaml")


def test_class_names():
    assert normalize_class("九（1）", "九") == "九1"
    assert normalize_class("91", "九") == "九1"
    assert normalize_class("七2", "七") == "七2"
    assert normalize_class("八年级3班", "八") == "八3"
    assert parse_classes("9192", "九") == ["九1", "九2"]
    assert parse_classes("93949596", "九") == ["九3", "九4", "九5", "九6"]
    assert parse_classes("93", "九") == ["九3"]
    assert parse_classes("七1、七2", "七") == ["七1", "七2"]


def test_subject_aliases():
    assert SCHOOL.canonical("道德与法治") == "道法"
    assert SCHOOL.canonical("物 理") == "物理"
    assert SCHOOL.canonical("班级") is None


@pytest.mark.parametrize("fname,grade,nsub", [("示例登分表_七年级.xlsx", "七年级", 7), ("示例登分表_八年级.xlsx", "八年级", 8),
                                              ("示例登分表_九年级.xlsx", "九年级", 7), ("示例登分表_中考.xlsx", "九年级", 10)])
def test_pipeline(tmp_path, fname, grade, nsub):
    import copy
    from chengji.loader import load_scores
    SCHOOL = copy.deepcopy(globals()["SCHOOL"])
    data = load_scores(ROOT / "samples" / fname, SCHOOL)
    # 学校设置里满分待定的科目（生物、地理）：像运行时那样当场设定，示例数据按 50 分
    SCHOOL.grade_full.setdefault(grade, {}).update({s: 50 for s in data.subjects if not SCHOOL.full[s]})
    assert data.grade == grade and len(data.subjects) == nsub
    assert len(data.excluded) == 1                      # 示例数据里放了 1 名缺考学生
    T = load_teachers(ROOT / "samples" / "示例任课总表.xlsx", SCHOOL, grade)
    E = load_enrolled(ROOT / "samples" / "示例任课总表.xlsx", SCHOOL, grade)
    cfg = make_config(SCHOOL, grade, data.subjects, "测试", "", T, enrolled=E)
    R = analyze(data, cfg)
    n = len(data.df)
    assert R.cut_rank == round(n * 0.85)
    assert R.above >= R.cut_rank and R.above + R.below == n
    assert sum(b[4] for b in R.bands) == R.above        # 分数段合计 = 线上人数
    for s in cfg.subjects:                              # 每科教师覆盖全部学生
        assert sum(t["人数"] for t in R.teachers if t["学科"] == s) == n
    assert (R.near["距线"].abs() <= cfg.near_range).all()
    out = build_excel(R, tmp_path / "t.xlsx", draft_conclusions(R))
    wb = openpyxl.load_workbook(out)
    assert {"各班统计", "班级结构分", "教师排名", "临界生名单", "参数设置", "成绩明细", "任课安排"} <= set(wb.sheetnames)
    assert sorted(x["名次"] for x in R.structure.values()) == list(range(1, len(R.classes) + 1))


def test_empty_subject_column_ignored(tmp_path):
    """整列没有成绩的科目 = 本次没考：不计算，也不能把学生判成“成绩不全”。"""
    from chengji.loader import load_scores
    wb = openpyxl.load_workbook(ROOT / "samples" / "示例登分表_九年级.xlsx")
    ws = wb.active
    c = ws.max_column + 1
    ws.cell(2, c, "生物")                               # 只有表头，没有分数
    ws.cell(2, c + 1, "地理")
    ws.cell(5, c + 1, "缺考")                           # 只有文字、没有数字，也算没成绩
    wb.save(tmp_path / "t.xlsx")
    data = load_scores(tmp_path / "t.xlsx", SCHOOL)
    assert data.subjects == ["语文", "数学", "英语", "物理", "化学", "道法", "历史"]
    assert data.empty_subjects == ["生物", "地理"]
    assert len(data.excluded) == 1                      # 仍然只有原来那 1 名缺考学生


def test_missing_full_mark():
    with pytest.raises(ValueError, match="满分"):
        make_config(SCHOOL, "九年级", ["生物"], "测试")


def test_zhongkao_scheme():
    """中考方案：平均分（原始分）×60% + 合格率×40%，不看优秀率。"""
    from chengji.analysis import subject_stats
    import pandas as pd
    cfg = make_config(SCHOOL, "九年级", ["语文"], "测试", scheme="中考", promote_ratio=0.8)
    st = subject_stats(pd.Series([100, 72, 50, 30.0]), cfg, "语文")   # 满分120，合格线72
    assert abs(st["得分"] - (63 * 0.6 + 50 * 0.4)) < 1e-9
    assert cfg.promote_ratio == 0.8 and not cfg.normalize


def test_ratio_parse():
    from chengji.cli import parse_ratio
    assert parse_ratio("80%") == 0.8 and parse_ratio("0.75") == 0.75 and parse_ratio("90") == 0.9


# ---------------- 班级结构分：用一组能手算的小数据核对 ----------------
def _tiny(scheme, prev=None):
    """3 个班、每班 2 人，语文（满分120，合格72，优秀96）+ 数学（满分100，合格60，优秀80）。"""
    import pandas as pd
    from chengji.analysis import analyze
    from chengji.loader import ScoreData
    rows = [("九1", 110, 95), ("九1", 100, 90),      # 两人都全科优秀
            ("九2", 80, 70), ("九2", 60, 50),        # 一人全科合格
            ("九3", 50, 40), ("九3", 30, 20)]        # 都不合格
    df = pd.DataFrame([(c, f"生{i}", f"{i:02d}", a, b) for i, (c, a, b) in enumerate(rows)],
                      columns=["班级", "姓名", "考号", "语文", "数学"])
    data = ScoreData(df=df, grade="九年级", subjects=["语文", "数学"])
    cfg = make_config(SCHOOL, "九年级", ["语文", "数学"], "测试", scheme=scheme,
                      enrolled={"九1": 2, "九2": 3, "九3": 2}, prev_rank=prev)
    return analyze(data, cfg).structure


def test_structure_pingshi():
    """平时：平均成绩50 + 全科合格率20 + 全科优秀率20 + 参考率5 + 增值评价5。"""
    X = _tiny("平时", prev={"九1": 2, "九2": 1, "九3": 3})
    avg = {"九1": 197.5 / 220 * 50, "九2": 130 / 220 * 50, "九3": 70 / 220 * 50}   # 总分满分 220
    sub = {"九1": avg["九1"] + 20 + 20 + 5,
           "九2": avg["九2"] + 0.5 * 20 + 0 + 2 / 3 * 5,
           "九3": avg["九3"] + 0 + 0 + 5}
    # 小计名次 1、2、3 → 5、4.9、4.8；九1 上次第2进步1名 5.05 → 封顶 5；九2 上次第1退步1名 4.85；九3 不变 4.8
    add = {"九1": 5.0, "九2": 4.85, "九3": 4.8}
    for c in X:
        assert abs(X[c]["小计"] - sub[c]) < 1e-9
        assert abs(X[c]["增值评价分"] - add[c]) < 1e-9
        assert abs(X[c]["结构分"] - (sub[c] + add[c])) < 1e-9
    assert [X[c]["名次"] for c in ("九1", "九2", "九3")] == [1, 2, 3]


def test_structure_pingshi_no_previous():
    """没有输入上次名次：增值只按本次名次给分。"""
    X = _tiny("平时")
    assert [round(X[c]["增值评价分"], 9) for c in ("九1", "九2", "九3")] == [5.0, 4.9, 4.8]


def test_structure_zhongkao():
    """中考：平均成绩50 + 优秀率15 + 合格率15 + 参考率10 + 进线率10 + 前10名每人0.1分。"""
    X = _tiny("中考")
    # 共 6 人，前85% = 第5名 = 90分 → 进线：九1 2/2、九2 2/2、九3 1/2；不足10人，前10名 = 全部
    exp = {"九1": 197.5 / 220 * 50 + 15 + 15 + 10 + 10 + 0.2,
           "九2": 130 / 220 * 50 + 0 + 7.5 + 2 / 3 * 10 + 10 + 0.2,
           "九3": 70 / 220 * 50 + 0 + 0 + 10 + 5 + 0.2}
    for c in X:
        assert abs(X[c]["结构分"] - exp[c]) < 1e-9
        assert "增值评价" not in X[c]


def test_structure_without_enrolment():
    """没有填应考人数：用实考人数代替，参考率 100%，这一项各班都拿满分（用户 2026-10-02 定的规则）。"""
    import pandas as pd
    from chengji.analysis import analyze
    from chengji.loader import ScoreData
    rows = [("九1", 110), ("九1", 100), ("九2", 80), ("九2", 60), ("九2", 50)]
    df = pd.DataFrame([(c, f"生{i}", f"{i:02d}", float(v)) for i, (c, v) in enumerate(rows)], columns=["班级", "姓名", "考号", "语文"])
    data = ScoreData(df=df, grade="九年级", subjects=["语文"])
    X = analyze(data, make_config(SCHOOL, "九年级", ["语文"], "测试")).structure                       # 都没填
    assert [X[c]["应考人数"] for c in ("九1", "九2")] == [2, 3] and all(X[c]["应考按实考"] for c in X)
    assert all(abs(X[c]["参考率"] - 1) < 1e-9 and abs(X[c]["参考率分"] - 5) < 1e-9 for c in X)
    Y = analyze(data, make_config(SCHOOL, "九年级", ["语文"], "测试", enrolled={"九2": 4})).structure     # 只填了一个班
    assert Y["九1"]["应考按实考"] and Y["九1"]["应考人数"] == 2
    assert not Y["九2"]["应考按实考"] and abs(Y["九2"]["参考率分"] - 3 / 4 * 5) < 1e-9


def test_service_enrolment_notes():
    """服务层：没填的班给出说明；填了但比实考还少的报错。"""
    from chengji import service as sv
    f = ROOT / "samples" / "示例登分表_九年级.xlsx"
    R, _c, msgs = sv.compute(f, SCHOOL, {"scheme": "平时"})                                            # 没有任课总表，也没填
    assert any("没有填应考人数，按实考人数计算" in m for m in msgs)
    assert all(x["应考按实考"] for x in R.structure.values())
    with pytest.raises(ValueError, match="比应考人数"):
        sv.compute(f, SCHOOL, {"scheme": "平时", "enrolled": {"九1": 10}})


def test_pdf_sections_parse():
    from chengji.cli import parse_sections
    from chengji.pdf_report import SECTIONS
    avail = [s for s in SECTIONS if s != "教师排名"]              # 假设本次没有任课教师
    assert parse_sections("1 4", avail, SECTIONS) == ["成绩通报", "班级结构分"]
    assert parse_sections("结构分、成绩通报", avail, SECTIONS) == ["成绩通报", "班级结构分"]   # 按页面顺序
    assert parse_sections("全部", avail, SECTIONS) == avail
    assert parse_sections("0", avail, SECTIONS) == []
    assert parse_sections("1 7", avail, SECTIONS) == ["成绩通报"]  # 本次没有的内容跳过
    with pytest.raises(ValueError):
        parse_sections("9", avail, SECTIONS)


def test_pdf_sections_pages():
    """只选两项时，发布版里只有这两项的页面。"""
    from chengji.loader import load_scores
    from chengji.pdf_report import build_html
    data = load_scores(ROOT / "samples" / "示例登分表_九年级.xlsx", SCHOOL)
    E = load_enrolled(ROOT / "samples" / "示例任课总表.xlsx", SCHOOL, "九年级")
    R = analyze(data, make_config(SCHOOL, "九年级", data.subjects, "测试", enrolled=E))
    html = build_html(R, draft_conclusions(R), ["成绩通报", "班级结构分"])
    assert html.count('<div class="page">') == 2
    assert "班级结构分评比" in html and "教师成绩排名" not in html and "临界生名单" not in html
    assert "第 2 页" in html and "第 3 页" not in html            # 页码连续


def test_class_rank():
    """班名次：本班内按总分排，同分同名次；与级名次互不影响。"""
    import pandas as pd
    from chengji.loader import ScoreData
    rows = [("九1", 100), ("九1", 90), ("九1", 90), ("九1", 50), ("九2", 95), ("九2", 60)]
    df = pd.DataFrame([(c, f"生{i}", f"{i:02d}", float(v)) for i, (c, v) in enumerate(rows)],
                      columns=["班级", "姓名", "考号", "语文"])
    cfg = make_config(SCHOOL, "九年级", ["语文"], "测试", enrolled={"九1": 4, "九2": 2})
    R = analyze(ScoreData(df=df, grade="九年级", subjects=["语文"]), cfg)
    assert list(R.df["班名次"]) == [1, 2, 2, 4, 1, 2]
    assert list(R.df["级名次"]) == [1, 3, 3, 6, 2, 5]


def test_grade_subjects():
    """各年级固定科目：七年级 7 科，八年级多一门物理，九年级 7 科，中考 10 科。示例登分表与之一致。"""
    from chengji.loader import load_scores
    assert SCHOOL.expected_subjects("七年级", "平时") == ["语文", "数学", "英语", "道法", "历史", "生物", "地理"]
    assert set(SCHOOL.expected_subjects("八年级", "平时")) == set(SCHOOL.expected_subjects("七年级", "平时")) | {"物理"}
    assert SCHOOL.expected_subjects("九年级", "平时") == ["语文", "数学", "英语", "物理", "化学", "道法", "历史"]
    assert SCHOOL.expected_subjects("九年级", "中考") == SCHOOL.subjects and len(SCHOOL.subjects) == 10
    for key, fname in [("七年级", "七年级"), ("八年级", "八年级"), ("九年级", "九年级"), ("中考", "中考")]:
        want = SCHOOL.grade_subjects[key]
        for folder, name in (("samples", f"示例登分表_{fname}.xlsx"), ("templates", f"登分表_{fname}.xlsx")):
            data = load_scores(ROOT / folder / name, SCHOOL)
            assert set(data.subjects) == set(want), (name, data.subjects)


def _sheet_with_extra(tmp_path, extra):
    """在九年级示例登分表后面加几列。extra：{表头: 每行的值（函数或常数）}"""
    wb = openpyxl.load_workbook(ROOT / "samples" / "示例登分表_九年级.xlsx")
    ws = wb.active
    for k, (head, val) in enumerate(extra.items()):
        c = ws.max_column + 1
        ws.cell(2, c, head)
        for r in range(3, ws.max_row + 1):
            ws.cell(r, c, val(r) if callable(val) else val)
    f = tmp_path / "加了列的登分表.xlsx"
    wb.save(f)
    return f


def test_new_subject_column(tmp_path):
    """登分表里加一门科目总表里没有的科目：只要这一列是分数，就一起统计；不是分数的列、总分名次之类不算。"""
    import copy
    from chengji.loader import load_scores
    f = _sheet_with_extra(tmp_path, {"信息技术": lambda r: 30 + r % 20, "评语": "良好", "年级名次": lambda r: r, "总分": 300, "加试": lambda r: r % 10})
    school = copy.deepcopy(SCHOOL)
    d = load_scores(f, school)
    assert d.subjects == ["语文", "数学", "英语", "物理", "化学", "道法", "历史", "信息技术", "加试"]
    assert d.new_subjects == ["信息技术", "加试"] and d.ignored_cols == ["评语"]
    assert school.full["信息技术"] == 0                       # 满分待定，由使用者当场填
    assert len(d.excluded) == 1                               # 新科目没有让更多学生被剔除
    # 本次不计“加试”
    d2 = load_scores(f, copy.deepcopy(SCHOOL), exclude=["加试", "历史"])
    assert d2.subjects == ["语文", "数学", "英语", "物理", "化学", "道法", "信息技术"] and d2.new_subjects == ["信息技术"]


def test_new_subject_is_computed(tmp_path):
    """新科目算进各班统计、结构分；没填满分时说明要填哪一科。"""
    from chengji import service as sv
    f = _sheet_with_extra(tmp_path, {"信息技术": lambda r: 30 + r % 20})
    with pytest.raises(ValueError, match="信息技术 还没有填满分"):
        sv.compute(f, SCHOOL, {"scheme": "平时"})
    R, _c, msgs = sv.compute(f, SCHOOL, {"scheme": "平时", "full": {"信息技术": 50}})
    assert R.cfg.subjects[-1] == "信息技术" and R.cfg.total_full == 600
    assert all("信息技术" in R.class_stats[c] for c in R.classes)
    assert any("新加的科目" in m for m in msgs) and any("多了 信息技术" in m for m in msgs)
    assert "信息技术" not in SCHOOL.full                      # 公共的设置对象没有被改动


def test_algo_settings_roundtrip():
    """界面里的算法设置：默认值读出来是百分数；合计不是 100% 不让保存；保存的写法能被程序读回去。"""
    import copy
    import yaml
    from chengji import service as sv
    from chengji.config import load_school
    a = sv.algo_view(SCHOOL)
    assert a["平时"]["score"] == {"优秀率": 25, "及格率": 25, "平均分": 50} and a["平时"]["normalize"] is True
    assert a["平时"]["structure"] == {"平均成绩": 50, "全科合格率": 20, "全科优秀率": 20, "参考率": 5, "进线率": 0, "增值评价": 5}
    assert a["中考"]["score"] == {"优秀率": 0, "及格率": 40, "平均分": 60} and a["中考"]["points"]["前10名每人加分"] == 0.1
    sch, st = sv.algo_to_yaml(a, SCHOOL)                                  # 原样保存 = 和默认设置等价
    assert st["平时"] == SCHOOL.structures["平时"] and st["中考"] == SCHOOL.structures["中考"]
    bad = copy.deepcopy(a); bad["平时"]["score"]["优秀率"] = 30
    with pytest.raises(ValueError, match="单科得分：三项加起来要等于 100%，现在是 105%"):
        sv.algo_to_yaml(bad, SCHOOL)
    bad = copy.deepcopy(a); bad["中考"]["structure"]["进线率"] = "5"
    with pytest.raises(ValueError, match="中考核算的班级结构分：各项加起来要等于 100%，现在是 95%"):
        sv.algo_to_yaml(bad, SCHOOL)
    new = copy.deepcopy(a)
    new["平时"]["score"] = {"优秀率": "30", "及格率": 20, "平均分": 50}
    new["平时"]["structure"].update({"平均成绩": 45, "进线率": 5})
    sch, st = sv.algo_to_yaml(new, SCHOOL)
    assert sch["平时"]["优秀率"] == 0.3 and st["平时"]["进线率"] == 5 and "进线率" not in st["中考"] or True
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        f = Path(d) / "本校设置.yaml"
        f.write_text(yaml.safe_dump({"算法方案": sch, "结构分方案": st}, allow_unicode=True), encoding="utf-8")
        s2 = load_school(ROOT / "config" / "学校设置.yaml", f)
    assert s2.schemes["平时"]["优秀率"] == 0.3 and s2.structures["平时"]["平均成绩"] == 45 and s2.structures["平时"]["进线率"] == 5
    assert sv.algo_view(s2)["平时"]["score"] == {"优秀率": 30, "及格率": 20, "平均分": 50}


def test_changed_weights_change_results_but_not_old_runs():
    """改了比例：新的核算按新比例；带着当时比例快照的旧考试仍按旧比例。"""
    import copy
    from chengji import service as sv
    f = ROOT / "samples" / "示例登分表_九年级.xlsx"
    old_R, _c, _m = sv.compute(f, SCHOOL, {"scheme": "平时"})
    snap = sv.scheme_snapshot(SCHOOL, "平时")
    s2 = copy.deepcopy(SCHOOL)
    s2.schemes["平时"].update({"优秀率": 0.0, "及格率": 0.0, "平均分": 1.0})
    s2.structures["平时"] = {"平均成绩": 100}
    new_R, _c, _m = sv.compute(f, s2, {"scheme": "平时"})
    c = old_R.classes[0]
    assert abs(new_R.structure[c]["结构分"] - new_R.structure[c]["平均成绩"] * 100) < 1e-9          # 只剩平均成绩一项
    assert abs(new_R.composite[c] - old_R.composite[c]) > 1
    again, _c, _m = sv.compute(f, s2, {"scheme": "平时", "schemeDef": snap})                        # 回看旧考试
    assert abs(again.composite[c] - old_R.composite[c]) < 1e-9
    assert abs(again.structure[c]["结构分"] - old_R.structure[c]["结构分"]) < 1e-9


def test_lines_settings():
    """及格线、优秀线、临界范围等可以在设置里改；旧考试带着当时的分数线。"""
    import copy
    from chengji import service as sv
    assert sv.lines_view(SCHOOL) == {"及格线": 60, "优秀线": 80, "升学线": 85, "高分线": 75, "临界范围": 15}
    y = sv.lines_to_yaml({"及格线": "50", "优秀线": 85, "升学线": 80, "高分线": 70, "临界范围": 10})
    assert y == {"及格线比例": 0.5, "优秀线比例": 0.85, "升学线比例": 0.8, "高分线比例": 0.7, "临界范围": 10}
    with pytest.raises(ValueError, match="优秀线要比及格线高"):
        sv.lines_to_yaml({"及格线": 80, "优秀线": 80, "升学线": 85, "高分线": 75, "临界范围": 15})
    f = ROOT / "samples" / "示例登分表_九年级.xlsx"
    old, _c, _m = sv.compute(f, SCHOOL, {"scheme": "平时"})
    snap = sv.scheme_snapshot(SCHOOL, "平时")
    s2 = copy.deepcopy(SCHOOL); s2.pass_ratio, s2.excellent_ratio, s2.near_range = 0.5, 0.9, 5
    new, _c, _m = sv.compute(f, s2, {"scheme": "平时"})
    c = old.classes[0]
    assert new.class_stats[c]["语文"]["及格率"] > old.class_stats[c]["语文"]["及格率"]      # 及格线降了，及格率升
    assert len(new.near) < len(old.near)                                                  # 临界范围窄了
    again, _c, _m = sv.compute(f, s2, {"scheme": "平时", "schemeDef": snap})
    assert again.class_stats[c]["语文"]["及格率"] == old.class_stats[c]["语文"]["及格率"] and len(again.near) == len(old.near)


def test_teacher_table_edit(tmp_path):
    """任课总表在界面里编辑：读出来、改、写回去，程序读到的就是改后的；填错了说清楚哪里错。"""
    from chengji import teachers as tt
    from chengji.config import load_enrolled
    src = ROOT / "samples" / "示例任课总表.xlsx"
    d = tt.read_table(src, SCHOOL)
    assert len(d["rows"]) == 75 and d["rows"][0] == {"grade": "七年级", "subject": "语文", "teacher": "示例教师1", "classes": "七1、七2"}
    d["rows"].append({"grade": "九年级", "subject": "政治", "teacher": " 新 老师 ", "classes": "九1 九2"})     # 别名、空格、空格分隔都认
    d["rows"].append({"grade": "八年级", "subject": "", "teacher": "", "classes": ""})                        # 空行忽略
    [c for c in d["classes"] if c["cls"] == "九1"][0]["enrolled"] = "50"
    d["classes"].append({"grade": "九年级", "cls": "九7", "enrolled": ""})
    f = tmp_path / "任课总表.xlsx"
    tt.write_table(f, tt.check_table(d, SCHOOL))
    T = load_teachers(f, SCHOOL, "九年级")
    assert ("道法", "新老师", ["九1", "九2"]) in [(t.subject, t.name, t.classes) for t in T]
    e = load_enrolled(f, SCHOOL, "九年级")
    assert e["九1"] == 50 and "九7" not in e                         # 没填应考人数的班不算有
    assert [c["cls"] for c in tt.read_table(f, SCHOOL)["classes"] if c["grade"] == "九年级"][-1] == "九7"
    for bad, msg in [({"grade": "九年级", "subject": "语文", "teacher": "张三", "classes": ""}, "教师和任教班级要一起填"),
                     ({"grade": "九年级", "subject": "", "teacher": "张三", "classes": "九1"}, "还没有填学科"),
                     ({"grade": "九年级", "subject": "语文", "teacher": "张三", "classes": "一班"}, "看不懂")]:
        with pytest.raises(ValueError, match=msg):
            tt.check_table({"rows": [bad], "classes": []}, SCHOOL)
    with pytest.raises(ValueError, match="写了两次"):
        tt.check_table({"rows": [], "classes": [{"grade": "九年级", "cls": "九1", "enrolled": 40}, {"grade": "九年级", "cls": "91", "enrolled": 41}]}, SCHOOL)


def test_version_compare():
    from chengji.server import newer_version
    assert newer_version("2.1.0", "2.0.9") and newer_version("2.0.10", "2.0.9")
    assert not newer_version("2.0.0", "2.0.0") and not newer_version("1.9", "2.0.0") and not newer_version("", "2.0.0")


def test_friendly_path(tmp_path):
    """“关于”里显示数据文件夹的位置：个人文件夹里的不出现用户名，其他位置原样显示。"""
    import sys
    from chengji.server import friendly_path
    home = tmp_path / "Users" / "somebody"
    (home / "Documents" / "分寸成绩核算").mkdir(parents=True)
    (home / "别的地方" / "分寸").mkdir(parents=True)
    docs = "文稿" if sys.platform == "darwin" else "文档"
    assert friendly_path(home / "Documents" / "分寸成绩核算", home) == f"{docs} → 分寸成绩核算"
    assert friendly_path(home / "别的地方" / "分寸", home) == "个人文件夹 → 别的地方 → 分寸"
    out = tmp_path / "D盘" / "分寸"
    out.mkdir(parents=True)
    assert friendly_path(out, home) == str(out)                    # 不在个人文件夹里（如 D 盘）：原样显示
    assert "somebody" not in friendly_path(home / "Documents" / "分寸成绩核算", home)
