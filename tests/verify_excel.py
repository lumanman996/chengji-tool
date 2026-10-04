"""核对：Excel 公式算出的结果 与 pandas 独立计算的结果 必须一致。
verify_compare() 另外核对“和上次比”的三张表（tests/verify_all.py 里调用）。

需要 LibreOffice 先把公式算出来（重算后另存），用法：
  python tests/verify_excel.py <登分表.xlsx> <已重算的统计表.xlsx> [--任课 data/任课总表.xlsx]
      [--方案 中考] [--升学比例 0.8] [--应考 "九1=46 九2=45"] [--上次名次 "九1=3 九2=1"] [--满分 "生物=50 地理=50"]
"""
import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from chengji.analysis import analyze  # noqa: E402
from chengji.cli import parse_class_values, parse_full_overrides  # noqa: E402
from chengji.config import load_enrolled, load_school, load_teachers, make_config  # noqa: E402
from chengji.loader import load_scores  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent


def verify(score_path, xlsx_path, teacher_path=None, scheme="平时", ratio=None, enrolled="", prev="", full="", school_path=None, local=None):
    school = load_school(school_path or ROOT / "config" / "学校设置.yaml", local)
    data = load_scores(score_path, school)
    school.grade_full.setdefault(data.grade, {}).update(parse_full_overrides(full))
    T = load_teachers(teacher_path, school, data.grade) if teacher_path else []
    gc = school.grades[data.grade]
    E = load_enrolled(teacher_path, school, data.grade) if teacher_path else {}
    E.update(parse_class_values(enrolled, gc))
    cfg = make_config(school, data.grade, data.subjects, "核对", "", T, scheme=scheme, promote_ratio=ratio,
                      enrolled=E, prev_rank=parse_class_values(prev, gc))
    R = analyze(data, cfg)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    errs, n = [], 0

    def eq(a, b, what):
        nonlocal n
        n += 1
        if a is None or abs(float(a) - float(b)) > 1e-6:
            errs.append(f"{what}: Excel={a} pandas={b}")

    s = wb["各班统计"]
    NS = len(cfg.subjects)
    for row in s.iter_rows(min_row=4, max_row=4 + len(R.classes), values_only=True):
        c = row[0]
        eq(row[1], R.online[c if c != "全年级" else "全年级"]["实考人数"], f"{c} 人数")
        for k, sub in enumerate(cfg.subjects):
            b = 2 + 5 * k
            st = R.class_stats[c][sub]
            for j, key in enumerate(["总分", "平均分", "及格率", "优秀率", "得分"]):
                eq(row[b + j], st[key], f"{c} {sub} {key}")
        eq(row[2 + 5 * NS], R.composite[c], f"{c} 综合")
        if c != "全年级":
            eq(row[3 + 5 * NS], R.class_rank[c], f"{c} 名次")
    z = wb[[n_ for n_ in wb.sheetnames if n_.endswith("分数段")][0]]
    eq(z["B4"].value, R.cut_rank, "升学线名次"); eq(z["B5"].value, R.cut, "升学线"); eq(z["B6"].value, R.above, "线上人数")
    cells = {str(r[0]): r for r in z.iter_rows(min_row=1, values_only=True) if r[0] is not None}
    for c in R.classes:
        r = cells[c]
        o = R.online[c]
        for j, key in [(2, "线上"), (3, "上线率"), (4, "线下"), (5, "线上临界"), (6, "线下临界"), (7, "高分")]:
            eq(r[j], o[key], f"{c} {key}")
    if "教师排名" in wb.sheetnames:
        xs = {(r[0], r[1]): r for r in wb["教师排名"].iter_rows(min_row=4, values_only=True) if r[0] in cfg.subjects}
        for t in R.teachers:
            r = xs[(t["学科"], t["教师"])]
            eq(r[3], t["人数"], f"{t['教师']} 人数"); eq(r[8], t["得分"], f"{t['教师']} 得分"); eq(r[9], t["名次"], f"{t['教师']} 名次")
    # 成绩明细：动态序号（未筛选时就是 1、2、3…）、总分、级名次、班名次
    d = wb["成绩明细"]
    hd = {c.value: j for j, c in enumerate(d[1])}
    drows = list(d.iter_rows(min_row=2, values_only=True))
    eq(len(drows), len(R.df), "成绩明细 行数")
    for i, (r, m) in enumerate(zip(drows, R.df.itertuples(index=False)), start=1):
        who = f"{m.班级}{m.姓名}"
        if r[hd["班级"]] != m.班级 or r[hd["姓名"]] != m.姓名:
            errs.append(f"成绩明细第{i}行 不是 {who}")
            continue
        eq(r[hd["序号"]], i, f"{who} 序号"); eq(r[hd["总分"]], m.总分, f"{who} 总分")
        eq(r[hd["级名次"]], m.级名次, f"{who} 级名次"); eq(r[hd["班名次"]], m.班名次, f"{who} 班名次")
    if R.structure:
        st = wb["班级结构分"]
        r3 = [c.value for c in st[3]]
        r4 = [c.value for c in st[4]]
        head = [b or a for a, b in zip(r3, r4)]                 # 合并的单列表头在第 3 行
        col = {h: j for j, h in enumerate(head) if h}
        keys = [k for k in ("应考人数", "实考人数", "小计", "小计名次", "增值评价分", "前10名人数", "加分",
                            "结构分", "名次") if k in col]
        keys += [k + "分" for k in ("平均成绩", "全科合格率", "全科优秀率", "参考率", "进线率") if k + "分" in col]
        rows = {r[0]: r for r in st.iter_rows(min_row=5, values_only=True) if r[0] in R.classes}
        eq(len(rows), len(R.classes), "结构分 班级数")
        for c in R.classes:
            for k in keys:
                eq(rows[c][col[k]], R.structure[c][k], f"{c} 结构分·{k}")
    q = wb["临界生名单"]
    nq = sum(1 for r in q.iter_rows(min_row=4, values_only=True) if r[0])
    eq(nq, len(R.near), "临界生人数")
    for r in q.iter_rows(min_row=4, values_only=True):
        if r[0]:
            m = R.near[(R.near["班级"] == r[0]) & (R.near["姓名"] == r[1])].iloc[0]
            n += 1
            if r[7 + NS] != m["薄弱学科"]:
                errs.append(f"{r[1]} 薄弱学科 Excel={r[7 + NS]} pandas={m['薄弱学科']}")
    return n, errs


def verify_compare(xlsx_path, R, C):
    """“和上次比”三张表：本次的公式结果要和 compare.build 用 pandas 结果算出来的一致；上次的固定值原样写入。"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    errs, n = [], 0

    def eq(a, b, what):
        nonlocal n
        n += 1
        if a is None or isinstance(a, str) or abs(float(a) - float(b)) > 1e-6:
            errs.append(f"{what}: Excel={a} pandas={b}")

    def same(a, b, what):
        nonlocal n
        n += 1
        if a != b:
            errs.append(f"{what}: Excel={a} 应为={b}")

    ws = wb["和上次比"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = {r[0]: i for i, r in enumerate(rows) if r[0] in ("班级", "科目")}
    col = lambda h: {v: j for j, v in enumerate(rows[hdr[h]]) if v}
    # 各班
    cc = col("班级")
    for i, x in enumerate(C["classes"]):
        r = rows[hdr["班级"] + 1 + i]
        same(r[0], x["name"], f"第{i + 1}个班")
        eq(r[cc["本次综合名次"]], R.class_rank[x["name"]], f"{x['name']} 本次综合名次")
        eq(r[cc["本次上线率"]], R.online[x["name"]]["上线率"], f"{x['name']} 本次上线率")
        if C["hasStructure"]:
            eq(r[cc["本次结构分"]], R.structure[x["name"]]["结构分"], f"{x['name']} 本次结构分")
            eq(r[cc["本次结构分名次"]], R.structure[x["name"]]["名次"], f"{x['name']} 本次结构分名次")
        if x["found"]:
            eq(r[cc["上次综合名次"]], x["prevRank"], f"{x['name']} 上次综合名次")
            eq(r[cc["本次综合名次"] + 1], x["prevRank"] - R.class_rank[x["name"]], f"{x['name']} 综合名次进退")
            eq(r[cc["本次上线率"] + 1], R.online[x["name"]]["上线率"] - x["prevOnlineRate"], f"{x['name']} 上线率变化")
            if C["hasStructure"]:
                eq(r[cc["本次结构分"] + 1], R.structure[x["name"]]["结构分"] - x["prevStruct"], f"{x['name']} 结构分变化")
                eq(r[cc["本次结构分名次"] + 1], x["prevStructRank"] - R.structure[x["name"]]["名次"], f"{x['name']} 结构分名次进退")
    # 各科
    sc = col("科目")
    g = R.class_stats["全年级"]
    for i, x in enumerate(C["subjects"]):
        r, s = rows[hdr["科目"] + 1 + i], x["subject"]
        same(r[0], s, f"第{i + 1}个科目")
        eq(r[sc["本次满分"]], R.cfg.full[s], f"{s} 本次满分")
        eq(r[sc["本次平均分"]], g[s]["平均分"], f"{s} 本次平均分")
        eq(r[sc["本次得分率"]], g[s]["平均分"] / R.cfg.full[s], f"{s} 本次得分率")
        eq(r[sc["本次得分率"] + 1], x["avgRate"] - x["prevAvgRate"], f"{s} 得分率变化")
        eq(r[sc["上次得分率"]], x["prevAvgRate"], f"{s} 上次得分率")
        eq(r[sc["本次及格率"]], g[s]["及格率"], f"{s} 本次及格率")
        eq(r[sc["本次及格率"] + 1], g[s]["及格率"] - x["prevPass"], f"{s} 及格率变化")
        eq(r[sc["本次优秀率"]], g[s]["优秀率"], f"{s} 本次优秀率")
        eq(r[sc["本次优秀率"] + 1], g[s]["优秀率"] - x["prevExc"], f"{s} 优秀率变化")
    for r in rows:
        if r[0] in ("上次线下临界生", "上次线上临界生"):
            t = C["nearUp"] if r[0] == "上次线下临界生" else C["nearDown"]
            eq(r[1], t["all"], f"{r[0]} 人数"); eq(r[2], t["found"], f"{r[0]} 能对上"); eq(r[4], t["moved"], f"{r[0]} {r[3]}")
    # 名次进退：每个人都按“班级 + 姓名”在 pandas 结果里找
    df = {(r.班级, r.姓名): {"总分": r.总分, "级名次": r.级名次} for r in R.df.itertuples(index=False)}   # 重名的不在对比名单里
    m = [r for r in wb["名次进退"].iter_rows(min_row=4, values_only=True) if r[0]]
    eq(len(m), len(C["students"]), "名次进退 人数")
    for r, x in zip(m, C["students"]):
        who = f"{r[0]}{r[1]}"
        same((r[0], r[1]), (x["cls"], x["name"]), f"名次进退 {who} 的位置")
        eq(r[3], df[(r[0], r[1])]["总分"], f"{who} 本次总分")
        eq(r[5], df[(r[0], r[1])]["级名次"], f"{who} 本次级名次")
        eq(r[6], x["prevRank"] - df[(r[0], r[1])]["级名次"], f"{who} 进退")
    # 上次临界生
    q = [r for r in wb["上次临界生"].iter_rows(min_row=4, values_only=True) if r[0]]
    eq(len(q), len(C["near"]), "上次临界生 人数")
    for r, x in zip(q, C["near"]):
        who = f"{r[0]}{r[1]}"
        same((r[0], r[1], r[2]), (x["cls"], x["name"], x["prevPos"]), f"上次临界生 {who}")
        if x["found"]:
            eq(r[4], df[(r[0], r[1])]["总分"], f"{who} 本次总分")
            same(r[5], "线上" if df[(r[0], r[1])]["总分"] >= R.cut else "线下", f"{who} 本次位置")
        same(r[6], x["where"], f"{who} 去向")
    return n, errs


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("scores"); ap.add_argument("xlsx"); ap.add_argument("--任课"); ap.add_argument("--方案", default="平时"); ap.add_argument("--升学比例", type=float)
    ap.add_argument("--应考", default=""); ap.add_argument("--上次名次", default=""); ap.add_argument("--满分", default="")
    a = ap.parse_args()
    n, errs = verify(a.scores, a.xlsx, a.任课, a.方案, a.升学比例, a.应考, a.上次名次, a.满分)
    print(f"核对 {n} 项，不一致 {len(errs)} 项")
    for e in errs[:30]:
        print("  ", e)
    sys.exit(1 if errs else 0)
